"""PICC record drafts and source-template-based chronological Word records."""
from copy import deepcopy
from datetime import date, datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image

from .db import audit, dump, now
from .domain import DomainError
from .identity import require_project

PICC_FORMAT = 'picc-record-v1'
TEMPLATE_PATH = Path(__file__).parent / 'templates' / 'picc-record-v1.docx'
FIELD_LIMITS = {'project_name': 120, 'survey_content': 200, 'location': 120,
                'survey_date': 10, 'record_number': 40, 'survey_notes': 400}
REQUIRED_FIELDS = {'project_name': '项目名称', 'survey_content': '查勘内容', 'location': '楼栋房号'}
TIME_LABELS = {'watermark': '水印时间', 'upload': '上传时间', 'missing': '时间异常'}
EDITABLE_FIELDS = ('project_name', 'survey_content', 'location')


def clean_text(value, limit):
    if not isinstance(value, str) or len(value) > limit or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value):
        raise DomainError(f'填写内容无效或超过{limit}字')
    return value.strip()


def normalize_time(value):
    if not isinstance(value, str) or not value.strip():
        return None
    try:
        # Date-only values must not silently acquire a fabricated midnight time.
        if len(value.strip()) < 16:
            return None
        parsed = datetime.fromisoformat(value.strip().replace('Z', '+00:00'))
        if parsed.tzinfo:
            parsed = parsed.astimezone(timezone(timedelta(hours=8)))
        return parsed.replace(tzinfo=None).isoformat()
    except ValueError:
        return None


def normalize_date(value):
    stamp = normalize_time(value)
    if stamp:
        return stamp[:10]
    try:
        return date.fromisoformat(value).isoformat() if isinstance(value, str) and len(value) == 10 else None
    except ValueError:
        return None


class PiccRecords:
    def __init__(self, db, domain):
        self.db, self.domain = db, domain

    def all_photos(self, case_id):
        self.domain.case(case_id)
        items, page = [], 1
        while True:
            result = self.domain.photos(case_id=case_id, page=page, page_size=500, sort='created_at')
            items.extend(result['items'])
            if len(items) >= result['total']:
                return items
            page += 1

    def survey_date(self, case_id, rows, photos):
        link = self.db.one('SELECT snapshot FROM platform_links WHERE case_id=?', (case_id,))
        if link:
            stamp = normalize_date(json.loads(link['snapshot']).get('survey', {}).get('siteSurveyDate'))
            if stamp:
                return stamp, '官网查勘信息'
        draft = self.db.one("SELECT d.values_json FROM platform_drafts d JOIN platform_links l "
                            "ON l.case_id=d.case_id AND l.report_receiving_id=d.target_report_id "
                            "WHERE d.case_id=? AND d.stage='survey' AND NOT EXISTS "
                            "(SELECT 1 FROM platform_draft_field_sources f WHERE f.case_id=d.case_id "
                            "AND f.stage='survey' AND f.field='siteSurveyDate' AND f.origin='automatic')", (case_id,))
        if draft:
            stamp = normalize_date(json.loads(draft['values_json']).get('siteSurveyDate'))
            if stamp:
                return stamp, '已填写的查勘信息'
        for visit in self.db.all("SELECT visit_date FROM visits WHERE case_id=? AND stage='survey' ORDER BY visit_date,id", (case_id,)):
            stamp = normalize_date(visit['visit_date'])
            if stamp:
                return stamp, '查勘到访记录'
        survey_ids = {p['id'] for p in rows if p.get('stage') == 'survey' or p.get('photo_group') in ('survey_checkin', 'survey_problem')}
        dated = [p for p in photos if p['sort_time']]
        survey = [p for p in dated if p.get('survey_photo') or p['id'] in survey_ids]
        candidates = survey or dated
        if candidates:
            first = candidates[0]
            return first['sort_time'][:10], ('最早查勘照片的' if survey else '最早照片的') + first['time_label']
        return '', '暂无查勘记录或照片时间'

    def get(self, case_id, rows=None):
        case = self.domain.case(case_id)
        draft = self.db.one('SELECT * FROM picc_record_drafts WHERE case_id=?', (case_id,))
        fields = {'project_name': case['project_name'], 'survey_content': case['title'],
                  'location': '-'.join(filter(None, (case.get('building'),
                               case.get('unit') + '单元' if case.get('unit') else '',
                               case.get('room') or case.get('location_detail')))),
                  'survey_date': '', 'record_number': '', 'survey_notes': ''}
        if draft:
            saved = json.loads(draft['fields'])
            fields.update({key: saved[key] for key in EDITABLE_FIELDS if key in saved})
        rows = self.all_photos(case_id) if rows is None else rows
        excluded_hashes = set(json.loads(draft['photo_overrides']).get('excluded_blob_hashes', [])) if draft else set()
        groups = {}
        for row in rows:
            if row['media_type'] == 'image':
                groups.setdefault(row['blob_hash'], []).append(row)
        # Prefer a valid watermark, then the original record over its work copy,
        # then the earliest upload. Filenames are never identity keys.
        unique = [min(values, key=lambda p: (not bool(normalize_time(p.get('watermark_time'))), bool(p.get('reuse_source_id')),
                                             p['created_at'], p['id'])) for values in groups.values()]
        photos = []
        for row in unique:
            sort_time, source = None, 'missing'
            for value, kind in ((row.get('watermark_time'), 'watermark'), (row.get('created_at'), 'upload')):
                parsed = normalize_time(value)
                if parsed:
                    sort_time, source = parsed, kind
                    break
            photos.append({'id': row['id'], 'original_name': row['original_name'],
                           'thumbnail_url': row['thumbnail_url'], 'sort_time': sort_time,
                           'time_source': source, 'time_label': TIME_LABELS[source],
                           'remarks': '', 'excluded': row['blob_hash'] in excluded_hashes,
                           'survey_photo': any(p.get('stage') == 'survey' or p.get('photo_group') in ('survey_checkin', 'survey_problem') for p in groups[row['blob_hash']]),
                           'duplicate_count': len(groups[row['blob_hash']]) - 1})
        photos.sort(key=lambda p: (p['sort_time'] or '9999', p['original_name'] or '', p['id']))
        fields['survey_date'], date_source = self.survey_date(case_id, rows, photos)
        excluded_photos = [p for p in photos if p['excluded']]
        photos = [p for p in photos if not p['excluded']]
        missing = sum(not p['sort_time'] for p in photos)
        blockers = ['请填写' + label for key, label in REQUIRED_FIELDS.items() if not fields[key]]
        if not case.get('platform_report_code'):
            blockers.append('请先在案件信息中补充报修编号')
        if not photos:
            blockers.append('本表照片已全部移除，请恢复至少一张照片' if excluded_photos else '本案尚无照片')
        elif not fields['survey_date']:
            blockers.append('暂无可用的查勘日期，请刷新案件信息或照片')
        if missing:
            blockers.append(f'{missing}张照片的上传时间异常，请刷新资料后重试')
        return {'case_id': case_id, 'revision': draft['revision'] if draft else 0,
                'fields': fields, 'survey_date_source': date_source, 'photos': photos, 'photo_overrides': {},
                'excluded_photos': excluded_photos, 'duplicate_count': sum(len(v) - 1 for v in groups.values()),
                'video_count': sum(p['media_type'] == 'video' for p in rows),
                'excluded_file_count': sum(p['media_type'] not in ('image', 'video') for p in rows),
                'missing_time_count': missing, 'blockers': blockers}

    def save(self, case_id, data):
        require_project(self.domain.case(case_id)['project_id'], manage=True)
        fields = data.get('fields')
        overrides = data.get('photo_overrides', {})
        if not isinstance(fields, dict) or not isinstance(overrides, dict) or set(fields) - FIELD_LIMITS.keys():
            raise DomainError('记录表填写格式无效')
        # Keep old client payloads compatible, but never reuse manual dates,
        # sort overrides or remarks from earlier drafts in this automatic plan.
        fields = {key: clean_text(fields.get(key, ''), FIELD_LIMITS[key]) for key in EDITABLE_FIELDS}
        action = data.get('photo_action')
        if action is not None and (not isinstance(action, dict) or not isinstance(action.get('excluded'), bool) or not isinstance(action.get('id'), str)):
            raise DomainError('照片移除操作格式无效')
        revision = data.get('revision')
        with self.db.tx() as con:
            before = con.execute('SELECT * FROM picc_record_drafts WHERE case_id=?', (case_id,)).fetchone()
            actual = before['revision'] if before else 0
            if isinstance(revision, bool) or revision != actual:
                raise DomainError('记录表已在其他页面更新，请刷新后重新填写', 409)
            excluded = set(json.loads(before['photo_overrides']).get('excluded_blob_hashes', [])) if before else set()
            if action is not None:
                photo = con.execute("SELECT p.blob_hash FROM photos p JOIN blobs b ON b.hash=p.blob_hash WHERE p.id=? AND p.case_id=? AND b.media_type='image'",
                                    (action['id'], case_id)).fetchone()
                if not photo:
                    raise DomainError('照片已变化或不属于本案，请刷新照片后重试', 409)
                if action['excluded']:
                    excluded.add(photo['blob_hash'])
                else:
                    excluded.discard(photo['blob_hash'])
            cleaned = {'excluded_blob_hashes': sorted(excluded)} if excluded else {}
            con.execute('INSERT INTO picc_record_drafts VALUES(?,?,?,?,?,?) ON CONFLICT(case_id) DO UPDATE SET '
                        'fields=excluded.fields,photo_overrides=excluded.photo_overrides,revision=excluded.revision,updated_at=excluded.updated_at',
                        (case_id, dump(fields), dump(cleaned), actual + 1, before['created_at'] if before else now(), now()))
            audit(con, 'picc_record', case_id, 'save', after={'revision': actual + 1, 'fields': fields, 'photo_overrides': cleaned})
        return self.get(case_id)

    def snapshot(self, case_id, rows):
        plan = self.get(case_id, rows)
        if plan['blockers']:
            raise DomainError('；'.join(plan['blockers']))
        return {'rules_version': 3, 'fields': plan['fields'], 'photos': [{key: p[key] for key in ('id', 'sort_time', 'time_source', 'remarks')} for p in plan['photos']],
                'template_sha256': hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest()}


def write_record(path, record, photos, folder):
    if hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest() != record['template_sha256']:
        raise DomainError('记录表模板已更新，请重新生成此版本')
    doc = Document(TEMPLATE_PATH)
    table = doc.tables[0]
    fields = record['fields']

    def fill(cell, text):
        paragraph = cell.paragraphs[0]
        props = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
        paragraph.clear()
        run = paragraph.add_run(text)
        if props is not None:
            run._r.insert(0, props)
        for extra in cell.paragraphs[1:]:
            cell._tc.remove(extra._p)

    for row, values in [(1, ('项目名称：' + fields['project_name'], '查勘内容：' + fields['survey_content'])),
                        (2, ('楼栋房号：' + fields['location'], '查勘日期：' + fields['survey_date'])),
                        (3, ('图片', '备注'))]:
        for cell, value in zip(table.rows[row].cells, values):
            fill(cell, value)
    prototype = deepcopy(table.rows[4]._tr)
    table._tbl.remove(table.rows[4]._tr)
    by_id = {p['id']: p for p in photos}
    for item in record['photos']:
        table._tbl.append(deepcopy(prototype))
        left, right = table.rows[-1].cells
        fill(left, '')
        paragraph = left.paragraphs[0]
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        photo = by_id[item['id']]
        image_path = folder / photo['package_path']
        with Image.open(image_path) as picture:
            width, height = picture.size
        scale = min(2.818 / width, 3.758 / height)
        paragraph.add_run().add_picture(str(image_path), width=Inches(width * scale), height=Inches(height * scale))
        fill(right, '')
        right.paragraphs[0].paragraph_format.line_spacing = 1.2
        for run in right.paragraphs[0].runs:
            run.font.size = Pt(11)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    # python-docx may normalize unrelated parts. Retain the reference-owned parts verbatim.
    mutable = {'word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'}
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    with zipfile.ZipFile(TEMPLATE_PATH) as archive:
        for name in archive.namelist():
            if name not in mutable:
                parts[name] = archive.read(name)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return {'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'size_bytes': path.stat().st_size}

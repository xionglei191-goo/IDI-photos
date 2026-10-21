from __future__ import annotations

import io
import json
import logging
import os
import re
import sqlite3
from contextlib import asynccontextmanager
from pathlib import Path
from urllib.parse import urlsplit

from fastapi import Body, FastAPI, File, Form, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response
from openpyxl import Workbook
from starlette.middleware.trustedhost import TrustedHostMiddleware

from . import __version__
from .archive_tree import ArchiveTree
from .config import ROOT, Settings
from .connectors.service import ConnectorService
from .db import Database, now, uid
from .domain import Domain, DomainError, PHOTO_GROUPS, ROLES, STAGES, STATUSES, WORKFLOW_STAGES
from .imports import Importer
from .jobs import JobRunner
from .platform import PlatformPreparation
from .platform_client import PlatformClient
from .platform_account import PlatformAccount
from .platform_routes import platform_routes
from .recognition import Recognizer
from .project_followup import ProjectFollowup
from .followup_routes import followup_routes
from .reports import Reporter
from .storage import Storage, safe_name
from .workflow import Workflow
from .workflow_routes import workflow_routes
from .review_batches import ReviewBatches
from .case_agent import CaseAgent
from .feishu_agent import FeishuAgent
from .feishu_agent_routes import feishu_agent_routes
from .agent_routes import agent_routes
from .workspace_chat import WorkspaceChat
from .workspace_chat_routes import workspace_chat_routes
from .photo_ai import PhotoAI
from .photo_workflow import PhotoWorkflow
from .automatic_uploads import AutomaticUploads, automatic_upload_routes
from .review_notifications import ReviewNotifications
from .operations import Operations, operations_routes
from .mobile_camera import MobileCamera, camera_routes
from .official_camera import OfficialCamera, official_camera_routes
from .photo_dedup import PhotoDedup
from .access import Access, access_routes
from .identity import acting_as, current
from .project_accounts import ProjectAccounts, project_account_routes


def create_app(settings: Settings | None = None):
    settings = settings or Settings.from_env()
    settings.prepare()
    db = Database(settings.data_dir / "archive.sqlite3")
    db.initialize()
    access = Access(db)
    domain = Domain(db)
    storage = Storage(settings, db)
    archive_tree = ArchiveTree(db, storage)
    importer = Importer(settings, db, storage, domain)
    recognizer = Recognizer(db, storage, domain)
    photo_dedup = PhotoDedup(db, storage, recognizer)
    importer.dedup = photo_dedup
    photo_dedup.backfill()
    reporter = Reporter(db, storage, domain)
    platform = PlatformPreparation(db, domain, storage, PlatformClient(write_enabled=settings.platform_write_enabled))
    project_accounts = ProjectAccounts(platform, settings.data_dir, write_enabled=settings.platform_write_enabled)
    platform.project_accounts = project_accounts
    platform.project_accounts.access = access
    platform.account = platform.project_accounts.company
    platform.client = platform.project_accounts.router
    jobs = JobRunner(db)
    workflow = Workflow(db, domain, platform, jobs)
    reviews = ReviewBatches(workflow)
    platform.workflow = workflow
    connectors = ConnectorService(settings, db, importer, jobs)
    followup = ProjectFollowup(settings, db, domain, storage, connectors, recognizer)
    photo_ai = PhotoAI(db, domain, storage, recognizer, followup)
    followup.photo_ai = photo_ai
    case_agent = CaseAgent(settings, db, domain, storage, recognizer, followup)
    followup.case_agent = case_agent
    feishu_agent = FeishuAgent(settings, db, domain, storage, recognizer, connectors, followup, platform, workflow, reviews, case_agent)
    feishu_agent.photo_ai = photo_ai
    photo_workflow = PhotoWorkflow(db, domain, storage, recognizer, platform, feishu_agent)
    feishu_agent.photo_workflow = photo_workflow
    automatic_uploads = AutomaticUploads(workflow, reviews, photo_workflow)
    reviews.automatic_uploads = automatic_uploads
    review_notifications = ReviewNotifications(automatic_uploads, followup)
    automatic_uploads.notifications = review_notifications
    workspace_chat = WorkspaceChat(feishu_agent)

    def import_job(payload, progress):
        result = importer.process(payload, progress)
        if payload.get("auto_ocr", True) and result.get("photo_ids") and db.setting("recognition", {}).get("local_ocr", True):
            pending = [photo_id for photo_id in result["photo_ids"] if not db.one("SELECT id FROM recognitions WHERE photo_id=? AND provider='local_ocr' AND state='completed'", (photo_id,))]
            if pending:
                result["recognition_job_id"] = jobs.enqueue("recognize", {"photo_ids": pending})["id"]
        return result

    jobs.register("import", import_job)
    jobs.register("ledger", importer.ledger_commit)
    jobs.register("recognize", recognizer.process)
    jobs.register("export", reporter.process)
    jobs.register("backup", reporter.backup)
    jobs.register("archive_tree", archive_tree.process)
    jobs.register("platform_sync", workflow.sync)

    @asynccontextmanager
    async def lifespan(app):
        with db.tx() as con:
            con.execute("UPDATE agent_maintenance_actions SET state='interrupted',error='服务中断；未自动重放，请查询当前状态后另行请求',updated_at=? WHERE state='running'", (now(),))
        if settings.worker_enabled:
            official_camera.start()
            project_accounts.start()
            jobs.start()
            workflow.synchronizer.start()
            connectors.start()
            archive_tree.start()
            followup.start()
            photo_ai.start()
            photo_workflow.start()
            automatic_uploads.start()
            review_notifications.start()
            case_agent.start()
            feishu_agent.start()
            workspace_chat.start()
        yield
        official_camera.stop()
        platform.client._closing.set()
        project_accounts.stop()
        workspace_chat.stop()
        review_notifications.stop()
        automatic_uploads.stop()
        photo_workflow.stop()
        photo_ai.stop()
        workflow.synchronizer.stop()
        platform.client._closing.set()
        feishu_agent.stop()
        case_agent.stop()
        followup.stop()
        connectors.stop()
        jobs.stop()
        archive_tree.stop()
        platform.client.close()

    app = FastAPI(title="IDI 资料管理", version=__version__, lifespan=lifespan, docs_url="/api/docs", redoc_url=None)
    app.state.services = {"settings": settings, "db": db, "domain": domain, "storage": storage, "importer": importer, "recognizer": recognizer, "reporter": reporter, "jobs": jobs, "connectors": connectors, "archive_tree": archive_tree, "platform": platform, "workflow": workflow, "reviews": reviews}
    app.state.services['followup'] = followup
    app.state.services['platform_account'] = platform.account
    app.state.services['project_accounts'] = platform.project_accounts
    app.include_router(project_account_routes(platform.project_accounts))
    app.state.services['photo_ai'] = photo_ai
    app.state.services['model_runtime'] = recognizer.runtime
    app.state.services['photo_workflow'] = photo_workflow
    app.state.services['automatic_uploads'] = automatic_uploads
    app.state.services['review_notifications'] = review_notifications
    app.include_router(automatic_upload_routes(automatic_uploads))
    app.state.services['case_agent'] = case_agent
    app.state.services['feishu_agent'] = feishu_agent
    app.state.services['workspace_chat'] = workspace_chat
    app.state.services['agent_memory'] = feishu_agent.memory
    operations = Operations(app.state.services)
    feishu_agent.operations = workspace_chat.operations = operations
    app.state.services['operations'] = operations
    app.include_router(operations_routes(operations))
    app.include_router(workspace_chat_routes(workspace_chat))
    app.include_router(feishu_agent_routes(feishu_agent))
    app.include_router(agent_routes(case_agent))
    app.include_router(followup_routes(followup))
    app.include_router(platform_routes(platform))
    app.include_router(workflow_routes(workflow, reviews))
    mobile_camera = MobileCamera(db, domain, storage)
    app.state.services['mobile_camera'] = mobile_camera
    app.include_router(camera_routes(mobile_camera))
    official_camera = OfficialCamera(mobile_camera)
    official_camera.dedup = photo_dedup
    app.state.services['official_camera'] = official_camera
    app.state.services['photo_dedup'] = photo_dedup
    app.state.services['access'] = access
    mobile_camera.access = access
    jobs.access = access
    app.include_router(access_routes(access))
    app.include_router(official_camera_routes(official_camera))
    app.add_middleware(TrustedHostMiddleware, allowed_hosts=settings.trusted_hosts)

    @app.middleware("http")
    async def network_boundary(request: Request, call_next):
        if request.client and not settings.allows_client(request.client.host):
            detail = "本版本仅允许本机访问" if settings.local_only else "仅允许本机及已配置局域网访问"
            return JSONResponse({"detail": detail}, status_code=403)
        host = request.headers.get("host", "")
        if host.split(":", 1)[0].lower() in settings.public_hosts:
            forwarded_origin = settings.proxy_origin(request.client.host if request.client else "", host,
                                                     request.headers.get("x-forwarded-proto", ""))
            if not forwarded_origin:
                return JSONResponse({"detail": "公网入口仅允许通过已配置的 HTTPS 隧道访问"}, status_code=403)
            # Trust only this proxy's scheme for this exact public hostname.
            # Keep the TCP peer address; never trust X-Forwarded-For for access.
            request.scope["scheme"] = "https"
        if request.url.path.startswith("/api/") and request.method not in ("GET", "HEAD", "OPTIONS"):
            if request.headers.get("X-IDI-Request") != "1":
                return JSONResponse({"detail": "缺少本地工作台请求标识"}, status_code=403)
            origin = request.headers.get("Origin")
            allowed_dev = os.environ.get("IDI_DEV_ORIGIN", "")
            if origin and (urlsplit(origin).netloc != request.headers.get("host") or urlsplit(origin).scheme != request.url.scheme) and origin != allowed_dev:
                return JSONResponse({"detail": "不允许来自其他网站的写操作"}, status_code=403)
        try:
            principal = await access.authorize(request)
            with acting_as(principal):
                await access.check_route(request)
                response = await call_next(request)
        except DomainError as error:
            response = JSONResponse({'detail': str(error)}, status_code=error.status)
            if error.status == 401:
                response.headers['X-IDI-Auth'] = 'expired'
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["Referrer-Policy"] = "same-origin"
        response.headers["X-Frame-Options"] = "SAMEORIGIN"
        if request.url.path.startswith("/api/"):
            response.headers["Cache-Control"] = "no-store"
        return response

    @app.exception_handler(DomainError)
    async def domain_error(request, error):
        return JSONResponse({"detail": str(error)}, status_code=error.status)

    @app.exception_handler(ValueError)
    async def invalid_value(request, error):
        return JSONResponse({"detail": str(error)}, status_code=400)

    @app.exception_handler(sqlite3.IntegrityError)
    async def integrity_error(request, error):
        logging.getLogger(__name__).warning("Database constraint rejected a request: %s", type(error).__name__)
        return JSONResponse({"detail": "数据关联冲突，请刷新后检查项目、案件或重复文件"}, status_code=409)

    def photo_ids(data):
        values = data.get("photo_ids", [])
        if not isinstance(values, list) or not values or len(values) > 500 or any(not isinstance(v, str) for v in values):
            raise DomainError("请选择1至500个文件")
        return list(dict.fromkeys(values))

    async def save_upload(file: UploadFile, limit=None):
        target = settings.data_dir / "staging" / (uid() + Path(file.filename or "file").suffix.lower())
        size = 0
        try:
            with target.open("wb") as out:
                while block := await file.read(1024 * 1024):
                    size += len(block)
                    if size > (limit or settings.max_file_bytes):
                        raise DomainError("上传文件超过大小限制")
                    out.write(block)
            return target, size
        except BaseException:
            target.unlink(missing_ok=True)
            raise
        finally:
            await file.close()

    @app.get("/api/health")
    def health():
        db.one("SELECT 1 AS ok")
        return {"status": "ok", "application": "IDI", "version": __version__, "time": now()}

    @app.get("/api/system")
    def system():
        if current() and not current().admin:
            return {"version": __version__, "stages": STAGES, "roles": ROLES, "photo_groups": PHOTO_GROUPS,
                    "case_statuses": STATUSES, "workflow_stages": WORKFLOW_STAGES,
                    "max_file_mb": settings.max_file_bytes // 1024 // 1024}
        return {"version": __version__, "data_dir": str(settings.data_dir), "local_only": settings.local_only, "deployment": settings.deployment,
                "listen_host": settings.host, "lan_urls": settings.lan_urls, "lan_subnets": list(settings.lan_subnets) if not settings.local_only else [],
                "public_urls": list(settings.public_origins), "trusted_proxies": list(settings.trusted_proxies),
                "stages": STAGES, "roles": ROLES, "photo_groups": PHOTO_GROUPS, "case_statuses": STATUSES, "workflow_stages": WORKFLOW_STAGES, "recognition": recognizer.status(),
                "max_file_mb": settings.max_file_bytes // 1024 // 1024, "worker_enabled": settings.worker_enabled}

    @app.get("/api/summary")
    def summary():
        return domain.summary()

    @app.get("/api/archive-tree")
    def archive_tree_status():
        return archive_tree.status()

    @app.post("/api/archive-tree/sync")
    def sync_archive_tree():
        return jobs.enqueue("archive_tree", {})

    @app.post("/api/archive-tree/open")
    def open_archive_tree():
        return archive_tree.open_folder()

    @app.get("/api/projects")
    def projects():
        return domain.projects()

    @app.post("/api/projects")
    def create_project(data: dict = Body(...)):
        return domain.save_project(data)

    @app.patch("/api/projects/{project_id}")
    def edit_project(project_id: str, data: dict = Body(...)):
        if not db.one("SELECT id FROM projects WHERE id=?", (project_id,)):
            raise DomainError("项目不存在", 404)
        return domain.save_project(data, project_id)

    @app.get("/api/cases")
    def cases(project_id: str | None = None, query: str | None = None, status: str | None = None):
        return domain.cases(project_id, query, status)

    @app.post("/api/cases")
    def create_case(data: dict = Body(...)):
        return domain.save_case(data)

    @app.get("/api/cases/{case_id}")
    def get_case(case_id: str):
        return domain.case(case_id)

    @app.patch("/api/cases/{case_id}")
    def edit_case(case_id: str, data: dict = Body(...)):
        return domain.save_case(data, case_id)

    @app.get("/api/cases/{case_id}/completeness")
    def completeness(case_id: str, template_id: str = "general"):
        return domain.completeness(case_id, template_id)

    @app.get("/api/cases/{case_id}/visits")
    def visits(case_id: str):
        return domain.visits(case_id)

    @app.post("/api/cases/{case_id}/visits")
    def add_visit(case_id: str, data: dict = Body(...)):
        return domain.save_visit(case_id, data)

    @app.put("/api/cases/{case_id}/visits/{visit_id}")
    def edit_visit(case_id: str, visit_id: str, data: dict = Body(...)):
        return domain.save_visit(case_id, data, visit_id)

    @app.get("/api/photos")
    def photos(project_id: str | None = None, case_id: str | None = None, stage: str | None = None, role: str | None = None,
               part: str | None = None, query: str | None = None, review: str | None = None, sort: str = "watermark_time", direction: str = "asc", page: int = 1, page_size: int = 30, photo_group: str | None = None, visit_id: str | None = None):
        result = domain.photos(project_id, case_id, stage, role, part, query, review, sort, direction, page, page_size, photo_group=photo_group, visit_id=visit_id)
        states = operations.photos([p['id'] for p in result['items']])
        for item in result['items']:
            item['processing'] = states[item['id']]
        return result

    @app.post("/api/photos/bulk")
    def bulk_photos(data: dict = Body(...)):
        ids = photo_ids(data)
        patch = data.get("patch", {})
        if not isinstance(patch, dict):
            raise DomainError("批量修改内容无效")
        results, errors = [], []
        for photo_id in ids:
            try:
                current = domain.update_photo(photo_id, patch)
                results.append(current["id"])
            except DomainError as error:
                errors.append({"photo_id": photo_id, "error": str(error)})
        return {"updated": results, "errors": errors, "state": "partial" if results and errors else "failed" if errors else "completed"}

    @app.get("/api/photos/{photo_id}")
    def photo(photo_id: str):
        return {**domain.photo(photo_id), 'processing': operations.photos([photo_id])[photo_id]}

    @app.post("/api/photos/{photo_id}/reuse-checkin")
    def reuse_checkin(photo_id: str, data: dict = Body(...)):
        if not storage.verify(domain.photo(photo_id)):
            raise DomainError('查勘原照缺失或校验失败，请先核对原件', 409)
        with platform.material_lock:
            return domain.reuse_checkin(photo_id, data)

    @app.post('/api/photos/delete-preview')
    def preview_photo_deletion(data: dict = Body(...)):
        from .photo_deletion import preview
        return preview(domain, photo_ids(data))

    @app.post('/api/photos/delete')
    def remove_photos(data: dict = Body(...)):
        from .photo_deletion import delete_photos
        # Match rule processing's lock order and serialize against official writes
        # and follow-up review, so none can persist stale photo evidence.
        with followup._send_lock, followup._review_lock, photo_workflow._lock, platform.material_lock:
            return delete_photos(domain, photo_ids(data), data.get('revisions'))

    @app.delete('/api/photos/{photo_id}')
    def remove_photo(photo_id: str, revision: int | None = None):
        return remove_photos({'photo_ids': [photo_id], 'revisions': {photo_id: revision} if revision is not None else None})

    @app.patch("/api/photos/{photo_id}")
    def edit_photo(photo_id: str, data: dict = Body(...)):
        return domain.update_photo(photo_id, data)

    @app.get("/api/photos/{photo_id}/original")
    def original(photo_id: str, download: bool = False):
        photo = domain.photo(photo_id)
        blob = db.one("SELECT * FROM blobs WHERE hash=?", (photo["blob_hash"],))
        path = storage.resolve(blob["stored_path"])
        if not path.is_file():
            raise DomainError("原文件丢失，请检查档案或备份", 404)
        filename = safe_name(photo.get("original_name") or path.name, 180)
        return FileResponse(path, media_type=blob["mime_type"], filename=filename if download else None, content_disposition_type="attachment" if download else "inline")

    @app.get("/api/photos/{photo_id}/thumbnail")
    def thumbnail(photo_id: str):
        row = db.one("SELECT b.thumbnail_path FROM photos p JOIN blobs b ON b.hash=p.blob_hash WHERE p.id=?", (photo_id,))
        if not row or not row["thumbnail_path"] or not storage.resolve(row["thumbnail_path"]).is_file():
            raise DomainError("该文件没有缩略图", 404)
        return FileResponse(storage.resolve(row["thumbnail_path"]), media_type="image/jpeg", headers={"Cache-Control": "private,max-age=86400"})

    @app.post("/api/import/folder")
    def import_folder(data: dict = Body(...)):
        entries = importer.folder_entries(str(data.get("path", "")))
        if not entries:
            raise DomainError("文件夹中没有支持的文件")
        batch_id = importer.create_batch("folder", data.get("project_id"), data.get("case_id"), {"folder": str(data["path"])})
        job = jobs.enqueue("import", {"batch_id": batch_id, "folder": str(Path(data["path"]).resolve()), "auto_ocr": data.get("auto_ocr", True)})
        return {"batch_id": batch_id, "job": job, "file_count": len(entries)}

    @app.post("/api/import/upload")
    async def import_upload(files: list[UploadFile] = File(...), project_id: str = Form(""), case_id: str = Form(""), paths: str = Form("[]"), auto_ocr: bool = Form(True)):
        project_id = access.upload_scope(project_id or None, case_id or None)
        if not files or len(files) > settings.max_batch_files:
            raise DomainError("上传文件数量无效")
        try:
            names = json.loads(paths)
        except ValueError as error:
            raise DomainError("目录信息无效") from error
        if not isinstance(names, list) or any(not isinstance(x, str) for x in names):
            raise DomainError("目录信息无效")
        entries, total = [], 0
        for index, file in enumerate(files):
            original_name = names[index] if index < len(names) else file.filename or "未命名"
            path, size = await save_upload(file)
            total += size
            if total > settings.max_batch_bytes:
                raise DomainError("本批上传文件超过总大小限制")
            entries.append({"path": str(path), "name": original_name, "source_path": original_name})
        batch_id = importer.create_batch("upload", project_id or None, case_id or None)
        job = jobs.enqueue("import", {"batch_id": batch_id, "files": entries, "auto_ocr": auto_ocr})
        return {"batch_id": batch_id, "job": job, "file_count": len(entries)}

    @app.post("/api/import/zip")
    async def import_zip(file: UploadFile = File(...), project_id: str = Form(""), case_id: str = Form(""), auto_ocr: bool = Form(True)):
        name = file.filename or "资料.zip"
        path, _ = await save_upload(file, settings.max_batch_bytes)
        entries = importer.stage_zip(path)
        for entry in entries:
            entry["source_path"] = "zip:" + name + "/" + entry["name"]
        batch_id = importer.create_batch("zip", project_id or None, case_id or None, {"filename": name})
        job = jobs.enqueue("import", {"batch_id": batch_id, "files": entries, "auto_ocr": auto_ocr})
        return {"batch_id": batch_id, "job": job, "file_count": len(entries)}

    @app.post("/api/import/ledger-preview")
    async def ledger_preview(file: UploadFile = File(...), project_id: str = Form(...)):
        name = file.filename or "台账.xlsx"
        path, _ = await save_upload(file, 20 * 1024 * 1024)
        return importer.ledger_preview(path.read_bytes(), name, project_id)

    @app.post("/api/import/ledger-commit")
    def ledger_commit(data: dict = Body(...)):
        batch_id = str(data.get("batch_id", ""))
        batch = importer.batch(batch_id)
        if batch["kind"] != "ledger":
            raise DomainError("不是台账预览批次")
        active = db.one("SELECT id FROM jobs WHERE kind='ledger' AND state IN ('queued','running') AND json_extract(payload,'$.batch_id')=?", (batch_id,))
        if active:
            return {"batch_id": batch_id, "job": jobs.get(active["id"])}
        return {"batch_id": batch_id, "job": jobs.enqueue("ledger", {"batch_id": batch_id})}

    @app.get("/api/ledger-template")
    def ledger_template():
        book = Workbook()
        sheet = book.active
        sheet.title = "案件台账"
        sheet.append(["报案时间", "地址及问题描述", "报案号", "栋号", "单元", "楼层", "房号", "位置类型", "详细位置", "资料提交时间", "报价时间", "保险公司", "维修部位", "备注", "报修编号", "保单号", "负责人"])
        sheet.freeze_panes = "A2"
        for col in range(1, 18):
            sheet.column_dimensions[sheet.cell(1, col).column_letter].width = 24 if col != 2 else 48
            for row in range(2, 502):
                sheet.cell(row, col).number_format = "@"
        help_sheet = book.create_sheet("填写说明")
        for value in ["报案号、楼栋、房号均按文本保存，避免长编号丢失精度。", "案件内容必填；报修编号、保单号、报案号分别填写，短编号根据项目简称和楼栋房号自动生成。", "日期可用2026-10-15或2026-10-15 08:30；维修部位可用顿号分隔。", "位置类型可填户内或公共区域；照片通过资料导入功能单独接收。"]:
            help_sheet.append([value])
        help_sheet.column_dimensions["A"].width = 100
        stream = io.BytesIO()
        book.save(stream)
        return Response(stream.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=idi-ledger-template.xlsx"})

    @app.get("/api/imports")
    def imports():
        rows = db.all("SELECT * FROM import_batches ORDER BY created_at DESC LIMIT 100")
        for row in rows:
            details = json.loads(row["details"])
            row["details"] = {k: v for k, v in details.items() if k != "rows"}
        return rows

    @app.get("/api/imports/{batch_id}")
    def get_import(batch_id: str):
        return importer.batch(batch_id)

    @app.get("/api/jobs")
    def list_jobs():
        return jobs.list()

    @app.get("/api/jobs/{job_id}")
    def get_job(job_id: str):
        return jobs.get(job_id)

    @app.post("/api/jobs/{job_id}/retry")
    def retry_job(job_id: str):
        return jobs.retry(job_id)

    @app.post("/api/recognize")
    def recognize(data: dict = Body(...)):
        ids = photo_ids(data)
        if data.get("use_vision") and not recognizer.status()["vision_ready"]:
            raise DomainError("外部图片识别尚未配置并启用")
        return jobs.enqueue("recognize", {"photo_ids": ids, "use_vision": bool(data.get("use_vision")), "force": bool(data.get("force", True))})

    @app.get("/api/settings/recognition")
    def recognition_settings():
        return recognizer.status()

    @app.put("/api/settings/recognition")
    def edit_recognition(data: dict = Body(...)):
        return recognizer.save_config(data)

    @app.get("/api/templates")
    def templates():
        return domain.templates()

    @app.post("/api/templates")
    def create_template(data: dict = Body(...)):
        return domain.save_template(data)

    @app.put("/api/templates/{template_id}")
    def edit_template(template_id: str, data: dict = Body(...)):
        return domain.save_template(data, template_id)

    @app.get("/api/review")
    def review(case_id: str | None = None, template_id: str = "general"):
        cases = [domain.case(case_id)] if case_id else domain.cases()[:200]
        checks = [domain.completeness(case["id"], template_id) for case in cases]
        return {"items": checks, "case_count": len(cases), "summary": domain.summary()}

    @app.post("/api/packages/preview")
    def package_preview(data: dict = Body(...)):
        return reporter.preview(photo_ids(data), data.get("template_id", "general"))

    @app.post("/api/packages")
    def create_package(data: dict = Body(...)):
        ids = photo_ids(data)
        preview = reporter.preview(ids, data.get("template_id", "general"))
        if data.get("parent_id") and len(preview["cases"]) != 1:
            raise DomainError("补件请每次选择一个案件")
        package_ids = {g["case"]["id"]: uid() for g in preview["cases"]}
        return jobs.enqueue("export", {"photo_ids": ids, "template_id": data.get("template_id", "general"), "parent_id": data.get("parent_id"), "title": data.get("title", ""), "package_ids": package_ids, "groups": [{"case": {"id": g["case"]["id"]}, "photo_ids": g["photo_ids"]} for g in preview["cases"]]})

    @app.get("/api/packages")
    def packages(case_id: str | None = None):
        return reporter.packages(case_id)

    @app.get("/api/packages/{package_id}")
    def package(package_id: str):
        result = reporter.package(package_id)
        result.pop("zip_path", None)
        return result

    @app.get("/api/packages/{package_id}/download")
    def download_package(package_id: str):
        package = reporter.package(package_id)
        if package["state"] in ("building", "failed"):
            raise DomainError("资料包尚未生成完成", 409)
        path = storage.resolve(package["zip_path"])
        if not path.is_file():
            raise DomainError("报送文件缺失", 404)
        case = package["manifest"]["case"]
        filename = safe_name(case["project_name"] + "-" + case["title"], 100) + f"-v{package['version']}.zip"
        return FileResponse(path, media_type="application/zip", filename=filename)

    @app.post("/api/packages/{package_id}/events")
    def package_event(package_id: str, data: dict = Body(...)):
        return reporter.add_event(package_id, data)

    @app.post("/api/packages/{package_id}/receipt")
    async def receipt(package_id: str, state: str = Form(...), event_at: str = Form(""), note: str = Form(""), file: UploadFile | None = File(None)):
        attachment_hash = None
        if file:
            name = file.filename or "回执.pdf"
            path, _ = await save_upload(file)
            attachment_hash = storage.put(path.read_bytes(), name)["hash"]
        return reporter.add_event(package_id, {"state": state, "event_at": event_at, "note": note}, attachment_hash)

    @app.get("/api/receipts/{event_id}/attachment")
    def receipt_attachment(event_id: str):
        row = db.one("SELECT b.* FROM package_events e JOIN blobs b ON b.hash=e.attachment_hash WHERE e.id=?", (event_id,))
        if not row:
            raise DomainError("回执附件不存在", 404)
        return FileResponse(storage.resolve(row["stored_path"]), media_type=row["mime_type"], filename="回执" + row["extension"])

    @app.get("/api/connectors")
    def list_connectors():
        return connectors.list()

    @app.get("/api/connectors/{connector_id}")
    def get_connector(connector_id: str):
        return connectors.get(connector_id)

    @app.post("/api/connectors")
    def add_connector(data: dict = Body(...)):
        return connectors.save(data)

    @app.post('/api/connectors/feishu/probe')
    def probe_feishu(data: dict = Body(...)):
        return connectors.probe_feishu(data)

    @app.post('/api/connectors/{connector_id}/retry')
    def retry_feishu(connector_id: str):
        return connectors.retry_feishu(connector_id)

    @app.put("/api/connectors/{connector_id}")
    def update_connector(connector_id: str, data: dict = Body(...)):
        connectors.get(connector_id)
        return connectors.save(data, connector_id)

    @app.post("/api/connectors/{connector_id}/events")
    def connector_event(connector_id: str, data: dict = Body(...)):
        return connectors.ingest(connector_id, data)

    @app.post("/api/connectors/{connector_id}/start")
    def start_collector(connector_id: str):
        return connectors.start_collector(connector_id)

    @app.post("/api/connectors/{connector_id}/stop")
    def stop_collector(connector_id: str):
        connectors.stop_collector(connector_id)
        row = connectors.get(connector_id)
        row["enabled"] = False
        return connectors.save(row, connector_id)

    @app.get("/api/messages")
    def messages(connector_id: str | None = None, failed_only: bool = False):
        where, values = [], []
        if connector_id:
            where.append("connector_id=?")
            values.append(connector_id)
        if failed_only:
            where.append("media_state IN ('failed','pending')")
        rows = db.all("SELECT * FROM messages" + (" WHERE " + " AND ".join(where) if where else "") + " ORDER BY received_at DESC LIMIT 200", values)
        for row in rows:
            row["raw"] = json.loads(row["raw"])
        return rows

    @app.post("/api/messages/{message_id}/retry")
    def retry_message(message_id: str):
        return connectors.retry_message(message_id)

    @app.get("/api/audit")
    def audit_log(entity_type: str | None = None, entity_id: str | None = None):
        where, values = [], []
        for key, value in (("entity_type", entity_type), ("entity_id", entity_id)):
            if value:
                where.append(key + "=?")
                values.append(value)
        rows = db.all("SELECT * FROM audit" + (" WHERE " + " AND ".join(where) if where else "") + " ORDER BY created_at DESC LIMIT 200", values)
        for row in rows:
            for key in ("before_json", "after_json"):
                row[key] = json.loads(row[key]) if row[key] else None
        return rows

    @app.post("/api/backups")
    def backup():
        return jobs.enqueue("backup", {"backup_id": uid()})

    @app.get("/api/backups")
    def backups():
        return [{"id": p.stem.removeprefix("idi-backup-"), "filename": p.name, "size_bytes": p.stat().st_size, "modified_at": p.stat().st_mtime, "download_url": "/api/backups/" + p.stem.removeprefix("idi-backup-") + "/download"} for p in sorted((settings.data_dir / "backups").glob("idi-backup-*.zip"), key=lambda p: p.stat().st_mtime, reverse=True)]

    @app.get("/api/backups/{backup_id}/download")
    def download_backup(backup_id: str):
        if not re.fullmatch(r"[a-f0-9]{32}", backup_id):
            raise DomainError("备份标识无效")
        path = settings.data_dir / "backups" / ("idi-backup-" + backup_id + ".zip")
        if not path.is_file():
            raise DomainError("备份不存在", 404)
        return FileResponse(path, media_type="application/zip", filename=path.name)

    frontend = Path(os.environ.get("IDI_FRONTEND_DIR", ROOT / "frontend" / "dist"))

    @app.get("/{path:path}", include_in_schema=False)
    def serve_frontend(path: str):
        if path.startswith("api/"):
            raise DomainError("接口不存在", 404)
        target = (frontend / path).resolve()
        if not target.is_relative_to(frontend.resolve()):
            raise DomainError("无效路径", 404)
        if target.is_file():
            return FileResponse(target)
        index = frontend / "index.html"
        if index.is_file():
            return FileResponse(index, headers={"Cache-Control": "no-cache"})
        return JSONResponse({"detail": "前端尚未构建，请运行 scripts/setup.ps1"}, status_code=503)

    return app

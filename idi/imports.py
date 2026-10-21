from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path, PurePosixPath

from openpyxl import load_workbook

from .config import Settings
from .db import Database, audit, dump, now, rowdict, uid
from .domain import Domain, DomainError, photo_group
from .storage import ALLOWED_EXTS, Storage
from .identity import current, require_project


def infer_path_metadata(path: str):
    normalized = path.replace("\\", "/")
    stage = "acceptance" if "验收" in normalized or "完工" in normalized else "survey" if "查勘" in normalized or "勘查" in normalized else "work" if "施工" in normalized else "unknown"
    role = "unknown"
    if "打卡" in normalized:
        if "入户" in normalized or "门牌" in normalized:
            role = "door_checkin"
        elif "楼栋" in normalized and not any(s in normalized for s in ("房号", "户号", "门")):
            role = "building_checkin"
    elif stage == "survey":
        role = "problem"
    elif stage == "work":
        role = "construction"
    elif stage == "acceptance":
        role = "finished"
    return {"stage": stage, "role": role}


class Importer:
    def __init__(self, settings: Settings, db: Database, storage: Storage, domain: Domain):
        self.settings, self.db, self.storage, self.domain = settings, db, storage, domain
        self.dedup = None

    def create_batch(self, kind, project_id=None, case_id=None, details=None):
        if case_id:
            project_id = self.domain.case(case_id)["project_id"]
        elif project_id and not self.db.one("SELECT id FROM projects WHERE id=?", (project_id,)):
            raise DomainError("项目不存在")
        batch_id = uid()
        require_project(project_id)
        person = current()
        with self.db.tx() as con:
            con.execute("INSERT INTO import_batches(id,kind,project_id,case_id,details,created_at,user_id) VALUES(?,?,?,?,?,?,?)", (batch_id, kind, project_id, case_id, dump(details or {}), now(), person.id if person else None))
        return batch_id

    def batch(self, batch_id):
        result = self.db.one("SELECT * FROM import_batches WHERE id=?", (batch_id,))
        if not result:
            raise DomainError("导入批次不存在", 404)
        require_project(result['project_id'])
        result["details"] = json.loads(result["details"])
        result["items"] = self.db.all("SELECT * FROM import_items WHERE batch_id=? ORDER BY created_at", (batch_id,))
        return result

    def add_file(self, data: bytes, name: str, *, case_id=None, project_id=None, source_type="upload",
                 source_path="", batch_id=None, message=None, metadata=None, prepared_ocr=None):
        if case_id:
            project_id = self.domain.case(case_id)["project_id"]
        require_project(project_id)
        person = current()
        if not data or len(data) > self.settings.max_file_bytes:
            raise ValueError("文件为空或超过单文件大小限制")
        suffix = Path(name).suffix.lower()
        if suffix not in ALLOWED_EXTS:
            raise ValueError(f"不支持的文件类型：{suffix or '无扩展名'}")
        incoming_hash = hashlib.sha256(data).hexdigest()
        if prepared_ocr and prepared_ocr['blob_hash'] != incoming_hash:
            raise DomainError('识别后图片文件发生变化，请重新导入', 409)
        fields = infer_path_metadata(source_path or name)
        metadata = {**(metadata or {}), 'uploaded_by': person.id if person else None,
                    'uploaded_by_name': person.display_name if person else ''}
        for key in ("stage", "role", "part", "process", "watermark_time", "notes"):
            if key in metadata:
                fields[key] = metadata[key]
        message = message or {}
        source_key = hashlib.sha256(dump([source_path or name, incoming_hash, message.get("id", "")]).encode()).hexdigest()
        with self.db.tx() as con:
            if case_id:
                photo = rowdict(con.execute("SELECT * FROM photos WHERE case_id=? AND blob_hash=? AND reuse_source_id IS NULL", (case_id, incoming_hash)).fetchone())
            else:
                photo = rowdict(con.execute("SELECT * FROM photos WHERE case_id IS NULL AND COALESCE(project_id,'')=? AND blob_hash=? AND reuse_source_id IS NULL", (project_id or "", incoming_hash)).fetchone())
                if photo is None:
                    # A later folder copy must retain provenance on the already filed image.
                    # Ambiguous reuse across multiple cases remains unassigned for review.
                    candidates = con.execute("SELECT * FROM photos WHERE blob_hash=? AND reuse_source_id IS NULL" + (" AND project_id=?" if project_id else ''),
                                             (incoming_hash, project_id) if project_id else (incoming_hash,)).fetchall()
                    if len(candidates) == 1:
                        photo = dict(candidates[0])
            reason = ''
            if photo is None and source_type == 'feishu' and prepared_ocr and self.dedup:
                photo, reason = self.dedup.find(con, data, prepared_ocr['identity'],
                    case_id=prepared_ocr['case_id'], project_id=project_id,
                    scope_conflict=prepared_ocr['scope_conflict'])
            duplicate = photo is not None
            if duplicate:
                photo_id = photo["id"]
                blob = rowdict(con.execute('SELECT * FROM blobs WHERE hash=?', (photo['blob_hash'],)).fetchone())
                if blob['hash'] == incoming_hash:
                    # Preserve the existing restore/mismatch behavior for exact file retries.
                    blob = self.storage.put(data, name, connection=con)
                elif not self.storage.verify(blob):
                    raise DomainError('已有照片原件校验失败，未合并新来源', 409)
                if prepared_ocr and self.dedup and blob['hash'] == incoming_hash:
                    if not photo['watermark_code']:
                        self.dedup.remember(con, photo_id, prepared_ocr['identity'])
                    self.dedup.cache(con, prepared_ocr)
            else:
                blob = self.storage.put(data, name, connection=con)
                if blob["media_type"] != "image":
                    fields["role"] = "other"
                photo_id = uid()
                reviewed_metadata = any(k in metadata for k in ("stage", "role", "part", "process", "watermark_time"))
                con.execute("INSERT INTO photos(id,blob_hash,project_id,case_id,stage,role,part,process,watermark_time,exif_time,notes,recognition_state,created_at,updated_at,photo_group,uploaded_by) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (photo_id, blob["hash"], project_id, case_id, fields["stage"], fields["role"], fields.get("part", ""), fields.get("process", ""), fields.get("watermark_time"), blob.get("exif_time"), fields.get("notes", ""), "recognized" if reviewed_metadata else "pending", now(), now(), photo_group(fields['stage'], fields['role']), person.id if person else None))
                audit(con, "photo", photo_id, "import", after={"blob_hash": blob["hash"], "case_id": case_id, "source_type": source_type, "source_path": source_path, "metadata": metadata}, actor="import")
                if prepared_ocr and self.dedup:
                    self.dedup.remember(con, photo_id, prepared_ocr['identity'], reason)
                    self.dedup.cache(con, prepared_ocr)
                    if reason:
                        con.execute("UPDATE photos SET recognition_state='conflict' WHERE id=?", (photo_id,))
            if source_type == 'feishu':
                metadata = {**metadata, 'original_hash': incoming_hash}
                if prepared_ocr:
                    metadata.update(watermark_code=prepared_ocr['identity']['code'],
                                    watermark_confidence=prepared_ocr['identity']['confidence'],
                                    watermark_dedup_reason=reason)
                if duplicate:
                    metadata.update(deduplicated=True, primary_hash=blob['hash'])
            con.execute("INSERT OR IGNORE INTO photo_sources VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)", (uid(), photo_id, batch_id, message.get("id"), source_type, source_key, Path(name.replace("\\", "/")).name, source_path, message.get("group_name", ""), message.get("sender", ""), message.get("sent_at"), now(), dump(metadata)))
        return {"photo_id": photo_id, "duplicate": duplicate, "blob_hash": blob["hash"], "media_type": blob["media_type"],
                "incoming_hash": incoming_hash, "watermark_dedup_reason": reason}

    def folder_entries(self, folder: str):
        root = Path(folder).expanduser().resolve()
        if not root.is_dir():
            raise DomainError("导入文件夹不存在")
        if root.is_relative_to(self.settings.data_dir) or self.settings.data_dir.is_relative_to(root):
            raise DomainError("请选择原始资料目录，不能导入数据档案目录或其父目录")
        entries = []
        for path in root.rglob("*"):
            if not path.is_file():
                continue
            if not path.resolve().is_relative_to(root):
                raise DomainError("资料文件夹中存在指向目录外的链接")
            if path.suffix.lower() in ALLOWED_EXTS:
                entries.append({"path": str(path), "name": str(path.relative_to(root)), "source_path": str(path)})
                if len(entries) > self.settings.max_batch_files:
                    raise DomainError("文件夹文件数量超过单批限制")
        return sorted(entries, key=lambda e: e["name"])

    def _record_item(self, batch_id, name, status, photo_id=None, error=None):
        with self.db.tx() as con:
            old = con.execute("SELECT id FROM import_items WHERE batch_id=? AND source_name=?", (batch_id, name)).fetchone()
            if old:
                con.execute("UPDATE import_items SET status=?,photo_id=?,error=?,created_at=? WHERE id=?", (status, photo_id, error, now(), old["id"]))
            else:
                con.execute("INSERT INTO import_items VALUES(?,?,?,?,?,?,?)", (uid(), batch_id, name, status, photo_id, error, now()))

    def process(self, payload: dict, progress):
        batch_id = payload["batch_id"]
        batch = self.batch(batch_id)
        with self.db.tx() as con:
            con.execute("UPDATE import_batches SET state='running',finished_at=NULL WHERE id=?", (batch_id,))
        try:
            if payload.get("folder"):
                entries = self.folder_entries(payload["folder"])
            else:
                entries = payload.get("files", [])
            new_ids, total_bytes = [], 0
            if not entries:
                raise DomainError("未找到可导入的图片、视频或文档")
            if len(entries) > self.settings.max_batch_files:
                raise DomainError("单批文件数量超过限制")
            with self.db.tx() as con:
                con.execute("UPDATE import_batches SET total=? WHERE id=?", (len(entries), batch_id))
            name_counts = Counter(entry["name"] for entry in entries)
            for index, entry in enumerate(entries):
                name = entry["name"]
                item_name = f"{index + 1:05d} · {name}" if name_counts[name] > 1 else name
                done = self.db.one("SELECT status,photo_id FROM import_items WHERE batch_id=? AND source_name=?", (batch_id, item_name))
                if done and done["status"] in ("imported", "duplicate", "deleted"):
                    if done["photo_id"]:
                        new_ids.append(done["photo_id"])
                    progress(index + 1, len(entries))
                    continue
                try:
                    path = Path(entry["path"])
                    if path.stat().st_size > self.settings.max_file_bytes:
                        raise ValueError("文件超过单文件大小限制")
                    total_bytes += path.stat().st_size
                    if total_bytes > self.settings.max_batch_bytes:
                        raise ValueError("本批文件超过总大小限制")
                    result = self.add_file(path.read_bytes(), name, case_id=batch["case_id"], project_id=batch["project_id"], batch_id=batch_id,
                                           source_type=batch["kind"], source_path=entry.get("source_path", name), metadata=entry.get("metadata"))
                    self._record_item(batch_id, item_name, "duplicate" if result["duplicate"] else "imported", result["photo_id"])
                    if result["media_type"] == "image":
                        new_ids.append(result["photo_id"])
                except Exception as error:
                    self._record_item(batch_id, item_name, "failed", error=str(error)[:2000])
                progress(index + 1, len(entries))
            counts = {row["status"]: row["count"] for row in self.db.all("SELECT status,COUNT(*) AS count FROM import_items WHERE batch_id=? GROUP BY status", (batch_id,))}
            imported, duplicates, failed = (counts.get(x, 0) for x in ("imported", "duplicate", "failed"))
            state = "partial" if failed and imported + duplicates else "failed" if failed else "completed"
            with self.db.tx() as con:
                con.execute("UPDATE import_batches SET state=?,imported=?,duplicates=?,failed=?,finished_at=? WHERE id=?", (state, imported, duplicates, failed, now(), batch_id))
            return {"batch_id": batch_id, "state": state, "imported": imported, "duplicates": duplicates, "failed": failed, "photo_ids": list(dict.fromkeys(new_ids))}
        except Exception as error:
            with self.db.tx() as con:
                details = {**batch["details"], "error": str(error)[:2000]}
                con.execute("UPDATE import_batches SET state='failed',details=?,finished_at=? WHERE id=?", (dump(details), now(), batch_id))
            raise

    def stage_zip(self, archive_path: Path):
        staged = self.settings.data_dir / "staging" / uid()
        staged.mkdir(parents=True, exist_ok=True)
        entries, total = [], 0
        try:
            with zipfile.ZipFile(archive_path) as archive:
                files = [info for info in archive.infolist() if not info.is_dir()]
                if len(files) > self.settings.max_batch_files:
                    raise DomainError("ZIP文件数量超过限制")
                for info in files:
                    name = info.filename.replace("\\", "/")
                    parts = PurePosixPath(name)
                    if parts.is_absolute() or ".." in parts.parts or ":" in name or "\x00" in name:
                        raise DomainError("ZIP包含不安全路径")
                    if (info.external_attr >> 16) & 0o170000 == 0o120000:
                        raise DomainError("ZIP不能包含符号链接")
                    total += info.file_size
                    if info.file_size > self.settings.max_file_bytes or total > self.settings.max_batch_bytes:
                        raise DomainError("ZIP解压后大小超过限制")
                    if info.compress_size and info.file_size / info.compress_size > 2000:
                        raise DomainError("ZIP压缩比例异常")
                    if parts.suffix.lower() not in ALLOWED_EXTS:
                        continue
                    target = staged / (uid() + parts.suffix.lower())
                    with archive.open(info) as source, target.open("wb") as out:
                        size = 0
                        while block := source.read(1024 * 1024):
                            size += len(block)
                            if size > self.settings.max_file_bytes:
                                raise DomainError("ZIP文件大小超过限制")
                            out.write(block)
                    entries.append({"path": str(target), "name": name, "source_path": "zip:" + archive_path.name + "/" + name})
            if not entries:
                raise DomainError("ZIP中没有支持的文件")
            return entries
        except (zipfile.BadZipFile, RuntimeError) as error:
            raise DomainError("ZIP损坏或需要密码") from error

    def ledger_preview(self, data: bytes, filename: str, project_id: str):
        if not self.db.one("SELECT id FROM projects WHERE id=?", (project_id,)):
            raise DomainError("请先选择台账所属项目")
        suffix = Path(filename).suffix.lower()
        rows, numeric_long_cells = [], set()
        if suffix == ".csv":
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("gb18030")
            rows = list(csv.reader(io.StringIO(text)))
        elif suffix == ".xlsx":
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                sheet = workbook.active
                for row in sheet.iter_rows():
                    values = []
                    for cell in row:
                        value = cell.value
                        if cell.data_type == "f":
                            value = "[不接受公式单元格]"
                        elif isinstance(value, (int, float)) and abs(value) >= 10 ** 15:
                            numeric_long_cells.add((cell.row, cell.column))
                        values.append(value)
                    rows.append(values)
                    if len(rows) > 5001:
                        raise DomainError("台账超过5000行")
            finally:
                workbook.close()
        else:
            raise DomainError("台账请使用CSV或XLSX格式")
        if not rows:
            raise DomainError("台账为空")
        if len(rows) > 5001:
            raise DomainError("台账超过5000行")
        mapping = {"短编号": "short_code", "报修编号": "platform_report_code", "保单号": "policy_number", "保单编号": "policy_number", "报案编号": "report_number", "报案号": "report_number", "案件号": "report_number", "负责人": "responsible", "报案时间": "report_at", "地址及问题描述": "title", "案件内容": "title", "问题描述": "title", "栋号": "building", "楼栋": "building", "单元": "unit", "楼层": "floor", "房号": "room", "户号": "room", "位置类型": "location_type", "详细位置": "location_detail", "资料提交时间": "submitted_at", "报价时间": "quoted_at", "保险公司": "insurer", "维修部位": "repair_parts", "备注": "description"}
        header_index = next((i for i, row in enumerate(rows[:10]) if any(str(v or "").strip() in ("地址及问题描述", "案件内容", "问题描述") for v in row)), None)
        if header_index is None:
            raise DomainError("未找到“地址及问题描述”或“案件内容”列")
        headers = [str(value or "").strip() for value in rows[header_index]]
        parsed = []
        for index, row in enumerate(rows[header_index + 1:], header_index + 2):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            item = {"project_id": project_id}
            warnings, errors = [], []
            for column, header in enumerate(headers):
                if header not in mapping or column >= len(row):
                    continue
                value = row[column]
                if value is None:
                    continue
                if isinstance(value, (datetime, date)):
                    value = value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
                if isinstance(value, float) and value.is_integer():
                    value = str(int(value))
                value = str(value).strip()
                field = mapping[header]
                if value == "[不接受公式单元格]":
                    errors.append(header + "含公式，请转换为实际值")
                    continue
                if field in ("report_number", "platform_report_code", "policy_number", "short_code") and (index, column + 1) in numeric_long_cells:
                    errors.append("长编号使用数值单元格可能已丢失精度，请改为文本后重导")
                if field.endswith("_at"):
                    value = value.replace("/", "-")
                    try:
                        value = datetime.fromisoformat(value).isoformat(sep=" ")
                    except ValueError:
                        errors.append(header + "日期格式无效")
                if field == "repair_parts":
                    value = [v.strip() for v in re.split(r"[、,，;；]", value) if v.strip()]
                if field == "location_type":
                    value = "public" if value in ("公共区域", "公共", "public") else "household"
                item[field] = value
            if not item.get("title"):
                errors.append("案件内容为空")
            if not item.get("report_number"):
                warnings.append("暂无报案号，将按位置和内容建立临时案件")
            if any(t in item.get("title", "") + item.get("room", "") for t in ("楼梯间", "公共", "车库", "外立面", "屋面", "负二", "负一")) and "location_type" not in item:
                item["location_type"] = "public"
                warnings.append("疑似公共区域，导入后请核对位置类型")
            parsed.append({"row": index, "data": item, "warnings": warnings, "errors": errors})
        batch_id = self.create_batch("ledger", project_id, details={"filename": filename, "rows": parsed})
        with self.db.tx() as con:
            con.execute("UPDATE import_batches SET state='preview',total=? WHERE id=?", (len(parsed), batch_id))
        return {"batch_id": batch_id, "headers": headers, "rows": parsed, "valid": sum(not r["errors"] for r in parsed), "invalid": sum(bool(r["errors"]) for r in parsed)}

    def ledger_commit(self, payload, progress):
        batch_id = payload["batch_id"]
        batch = self.batch(batch_id)
        if batch["kind"] != "ledger":
            raise DomainError("不是台账导入批次")
        rows = batch["details"].get("rows", [])
        for index, row in enumerate(rows):
            name = f"第{row['row']}行"
            prior = self.db.one("SELECT status FROM import_items WHERE batch_id=? AND source_name=?", (batch_id, name))
            if prior and prior["status"] in ("imported", "duplicate"):
                progress(index + 1, len(rows))
                continue
            if row["errors"]:
                self._record_item(batch_id, name, "failed", error="；".join(row["errors"]))
            else:
                try:
                    value = row["data"]
                    if value.get("platform_report_code"):
                        found = self.db.one("SELECT id FROM cases WHERE platform_report_code=?", (value["platform_report_code"],))
                    elif value.get("short_code"):
                        found = self.db.one("SELECT id FROM cases WHERE short_code=?", (value["short_code"],))
                    else:
                        found = self.db.one("SELECT id FROM cases WHERE project_id=? AND title=? AND building=? AND room=?", (value["project_id"], value["title"], value.get("building", ""), value.get("room", "")))
                    if found:
                        self._record_item(batch_id, name, "duplicate", error="已有相同报修编号、短编号或同位置同内容记录，保留已有案件；如为另一案件请手动新建")
                    else:
                        created = self.domain.save_case(value)
                        self._record_item(batch_id, name, "imported")
                        with self.db.tx() as con:
                            audit(con, "case", created["id"], "ledger_import", after={"batch_id": batch_id, "row": row["row"], "values": value}, actor="import")
                except Exception as error:
                    self._record_item(batch_id, name, "failed", error=str(error))
            progress(index + 1, len(rows))
        counts = {r["status"]: r["count"] for r in self.db.all("SELECT status,COUNT(*) AS count FROM import_items WHERE batch_id=? GROUP BY status", (batch_id,))}
        imported, duplicates, failed = (counts.get(s, 0) for s in ("imported", "duplicate", "failed"))
        state = "partial" if failed and imported + duplicates else "failed" if failed else "completed"
        with self.db.tx() as con:
            con.execute("UPDATE import_batches SET state=?,imported=?,duplicates=?,failed=?,finished_at=? WHERE id=?", (state, imported, duplicates, failed, now(), batch_id))
        return {"batch_id": batch_id, "state": state, "imported": imported, "duplicates": duplicates, "failed": failed}

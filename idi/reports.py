from __future__ import annotations

import hashlib
import html
import json
import os
import shutil
import sqlite3
import threading
import zipfile
from collections import defaultdict
from contextlib import closing
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageOps

from .db import Database, audit, dump, now, uid
from .domain import Domain, DomainError, PHOTO_GROUPS, ROLES, STAGES, STATUSES, WORKFLOW_STAGES
from .storage import Storage, safe_name
from .package_content import content_fingerprint
from .photo_package import PHOTO_PACKAGE_FORMAT, PHOTO_DIRECTORIES, case_folder, photo_directory, photo_paths, write_media
from .picc_records import PICC_FORMAT, PiccRecords, write_record


def location(case):
    parts = [case.get("building", "") + ("栋" if case.get("building") else ""),
             case.get("unit", "") + ("单元" if case.get("unit") else ""), case.get("room", ""), case.get("location_detail", "")]
    return " ".join(p for p in parts if p) or "位置待补充"


def spreadsheet(path, manifest):
    book = Workbook()
    index = book.active
    index.title = "照片索引"
    headers = ["序号", "项目", "报案编号（保司）", "案件内容", "位置", "阶段", "用途", "维修部位", "施工内容", "水印时间", "EXIF时间", "消息时间", "原文件名", "文件SHA256", "来源数", "核对状态", "短编号", "报修编号（平台）", "保单号", "照片分类", "包内文件"]
    index.append(headers)
    case = manifest["case"]
    for number, photo in enumerate(manifest["photos"], 1):
        message_time = next((s["message_time"] for s in photo["sources"] if s["message_time"]), "")
        index.append([number, case["project_name"], case.get("report_number") or "", case["title"], location(case), STAGES[photo["stage"]], ROLES[photo["role"]], photo["part"], photo["process"], photo["watermark_time"], photo["exif_time"], message_time, photo["original_name"], photo["blob_hash"], len(photo["sources"]), "人工已确认" if photo["recognition_state"] == "confirmed" else "待人工核对", case.get("short_code"), case.get("platform_report_code"), case.get("policy_number"), PHOTO_GROUPS.get(photo.get("photo_group"), "待分类"), photo["package_path"]])
        index.cell(index.max_row, len(headers)).hyperlink = photo["package_path"]
    info = book.create_sheet("案件信息")
    info.append(["字段", "值"])
    for label, value in (("短编号", case.get("short_code") or "待补充"), ("报修编号（平台）", case.get("platform_report_code") or "待关联"), ("保单号", case.get("policy_number") or "待补充"), ("项目", case["project_name"]), ("报案编号（保司）", case.get("report_number") or "暂无"), ("案件内容", case["title"]), ("位置", location(case)), ("保险公司", case.get("insurer", "")), ("报案时间", case.get("report_at")), ("资料提交时间", case.get("submitted_at")), ("报价时间", case.get("quoted_at")), ("业务状态", WORKFLOW_STAGES.get(case.get("workflow_stage"), STATUSES[case["status"]])), ("快照生成时间", manifest["created_at"]), ("报送版本", manifest["version"]), ("资料要求", manifest["template"]["name"]), ("照片数量", len(manifest["photos"]))):
        info.append([label, value])
    missing = book.create_sheet("缺项核对")
    missing.append(["阶段", "用途", "范围", "要求数量", "本次数量", "缺少数量"])
    for row in manifest["completeness"]["items"]:
        missing.append([STAGES[row["stage"]], PHOTO_GROUPS.get(row.get("photo_group"), ROLES[row["role"]]), row["scope"], row["minimum"], row["count"], row["missing"]])
    sources = book.create_sheet("来源记录")
    sources.append(["包内文件", "原文件名", "来源类型", "群名", "发送人", "消息时间", "系统接收时间"])
    for photo in manifest["photos"]:
        for source in photo["sources"]:
            sources.append([photo["package_path"], source["original_name"], source["source_type"], source["group_name"], source["sender"], source["message_time"], source["received_at"]])
    if manifest.get('visits'):
        entries = book.create_sheet('进场维修记录')
        entries.append(['进场日期', '本次标记', '阶段', '开始时间', '结束时间', '人数', '维修记录', '本次资料', '共用查勘打卡'])
        names = {p['id']: p['original_name'] for p in manifest['photos']}
        for visit in manifest['visits']:
            entries.append([visit['visit_date'], visit['visit_key'], STAGES[visit['stage']], visit['started_at'], visit['finished_at'], visit['crew_number'], visit['notes'],
                            '\n'.join(names[pid] for pid in visit['photo_ids']), '\n'.join(names[pid] for pid in visit['shared_checkin_ids'])])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 28
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="244B6B")
            cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                cell.font = Font(name="微软雅黑", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.row_dimensions[row[0].row].height = 36
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 22
        sheet.sheet_view.showGridLines = False
    for column, width in {"A": 7, "D": 36, "I": 32, "M": 42, "N": 66, "Q": 54}.items():
        index.column_dimensions[column].width = width
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 70
    sources.column_dimensions["A"].width = 54
    sources.column_dimensions["B"].width = 48
    book.save(path)


def word_report(path, manifest, folder):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.65)
    section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.7)
    section.right_margin = Inches(.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.color.rgb = RGBColor.from_string("244B6B")
    doc.add_heading("IDI 案件照片资料报告", 0)
    case = manifest["case"]
    doc.add_paragraph(case["title"], "Subtitle")
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for label, value in (("短编号 / 报修编号", " / ".join(filter(None, (case.get("short_code"), case.get("platform_report_code")))) or "待补充"), ("保单号", case.get("policy_number") or "待补充"), ("项目 / 位置", case["project_name"] + " / " + location(case)), ("报案编号（保司）", case.get("report_number") or "暂无，按案件内容区分"), ("保险公司", case.get("insurer") or "待补充"), ("报案时间", case.get("report_at") or "待补充"), ("业务状态", WORKFLOW_STAGES.get(case.get("workflow_stage"), STATUSES[case["status"]])), ("报送快照", f"版本 {manifest['version']} · {manifest['created_at']}"), ("本次资料", f"{len(manifest['photos'])} 个文件")):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, str(value)
    doc.add_heading("资料核对", 1)
    missing = [r for r in manifest["completeness"]["items"] if r["missing"]]
    if missing:
        check = doc.add_table(rows=1, cols=4)
        check.style = "Light Shading Accent 1"
        for cell, value in zip(check.rows[0].cells, ["阶段 / 用途", "范围", "要求", "本次"]):
            cell.text = value
        for row in missing:
            for cell, value in zip(check.add_row().cells, [STAGES[row["stage"]] + " / " + PHOTO_GROUPS.get(row.get("photo_group"), ROLES[row["role"]]), row["scope"], str(row["minimum"]), str(row["count"])]):
                cell.text = value
    else:
        doc.add_paragraph("本次所选资料满足当前模板的数量要求。")
    doc.add_paragraph("本报告记录本次选材与资料核对情况；完工效果照片不等同于验收通过，报送包生成不等同于保险公司接收。日期字段分别保留水印、EXIF和消息来源，未提供的信息保持空白。")
    if manifest.get('visits'):
        doc.add_heading('进场与维修记录', 1)
        for visit in manifest['visits']:
            doc.add_heading(visit['visit_date'] + ' · ' + (visit['visit_key'] or STAGES[visit['stage']]), 2)
            doc.add_paragraph('开始：' + (visit['started_at'] or '未填') + '；结束：' + (visit['finished_at'] or '未填') + '；人数：' + str(visit['crew_number'] or '未填'))
            doc.add_paragraph(visit['notes'] or '未填写本次维修说明')
            doc.add_paragraph(f"本次选入资料 {len(visit['photo_ids'])} 项；共用查勘打卡 {len(visit['shared_checkin_ids'])} 张。原件未重复存放。")
    images = [p for p in manifest["photos"] if p["media_type"] == "image"]
    for index, photo in enumerate(images):
        if index % 2 == 0:
            doc.add_page_break()
            doc.add_heading("照片档案", 1)
        doc.add_heading(f"{index + 1:02d}  {STAGES[photo['stage']]} · {ROLES[photo['role']]} · {photo['part'] or '部位待补充'}", 2)
        source = folder / photo["package_path"]
        with Image.open(source) as image:
            image = ImageOps.exif_transpose(image).convert("RGB")
            width, height = image.size
            scale = min(6.2 / width, 2.7 / height)
            import io
            stream = io.BytesIO()
            image.thumbnail((1700, 1500))
            image.save(stream, "JPEG", quality=88)
            stream.seek(0)
            doc.add_picture(stream, width=Inches(width * scale), height=Inches(height * scale))
        doc.add_paragraph(f"施工内容：{photo['process'] or '未提供'}\n水印时间：{photo['watermark_time'] or '未提供'}   EXIF时间：{photo['exif_time'] or '未提供'}\n原文件名：{photo['original_name']}")
    footer = section.footer.paragraphs[0]
    footer.text = "IDI 案件资料 · 第 "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    doc.save(path)


def html_report(path, manifest):
    case = manifest["case"]
    esc = html.escape
    items = []
    for photo in manifest["photos"]:
        visual = f'<img src="{esc(photo["package_path"], quote=True)}" alt="{esc(photo["process"], quote=True)}">' if photo["media_type"] == "image" else f'<a href="{esc(photo["package_path"], quote=True)}">打开附件</a>'
        items.append(f'<article>{visual}<h2>{esc(STAGES[photo["stage"]])} · {esc(ROLES[photo["role"]])} · {esc(photo["part"])}</h2><p>{esc(photo["process"])}</p><p>水印时间：{esc(photo["watermark_time"] or "未提供")}</p><p>{esc(photo["original_name"])}</p></article>')
    path.write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>案件照片资料</title><style>body{font:15px/1.6 "Microsoft YaHei",sans-serif;margin:32px;color:#263847}h1{font-size:24px}h2{font-size:17px}article{break-inside:avoid;border-bottom:1px solid #ddd;padding:20px 0}img{max-width:100%;max-height:470px}p{overflow-wrap:anywhere}@media print{article{page-break-inside:avoid}img{max-height:320px}}</style><h1>'+esc(" · ".join(filter(None, (case.get("short_code"), case.get("platform_report_code"), case["title"]))))+"</h1><p>"+esc(case["project_name"]+" / "+location(case))+"</p><p>报案编号（保司）："+esc(case.get("report_number") or "暂无")+f" · 版本{manifest['version']}</p>"+"".join(items)+"</html>", encoding="utf-8")


class Reporter:
    def __init__(self, db: Database, storage: Storage, domain: Domain):
        self.db, self.storage, self.domain = db, storage, domain
        self.picc = PiccRecords(db, domain)
        self._package_lock = threading.RLock()

    def packages(self, case_id=None):
        from .identity import project_filter
        scope, args = project_filter('c.project_id')
        rows = self.db.all("SELECT pk.*,c.project_id,c.title AS case_title,pr.name AS project_name FROM packages pk JOIN cases c ON c.id=pk.case_id JOIN projects pr ON pr.id=c.project_id WHERE pk.state<>'deleted' AND " + scope + (" AND pk.case_id=?" if case_id else "") + " ORDER BY pk.created_at DESC,pk.version DESC", (*args, case_id) if case_id else args)
        for row in rows:
            manifest = json.loads(row.pop("manifest"))
            row["case_title"] = manifest["case"]["title"]
            row["project_name"] = manifest["case"]["project_name"]
            for field in ("short_code", "platform_report_code", "report_number", "building", "unit", "room", "location_detail"):
                row[field] = manifest["case"].get(field) or ""
            row.pop("zip_path", None)
            row["photo_count"] = len(manifest["photos"])
            row["missing_count"] = manifest["completeness"]["missing_count"]
            row["download_url"] = f"/api/packages/{row['id']}/download" if row["state"] in ("ready", "submitted", "accepted", "returned") else None
            row['document_url'] = f"/api/packages/{row['id']}/document" if row['download_url'] and manifest.get('document_path') else None
            row['format'] = manifest.get('format')
        return rows

    def package(self, package_id):
        row = self.db.one("SELECT * FROM packages WHERE id=?", (package_id,))
        if not row or row['state'] == 'deleted':
            raise DomainError("报送记录不存在", 404)
        self.domain.case(row['case_id'])
        row["manifest"] = json.loads(row["manifest"])
        row["events"] = self.db.all("SELECT * FROM package_events WHERE package_id=? ORDER BY created_at", (package_id,))
        row["download_url"] = f"/api/packages/{package_id}/download" if row['state'] in ('ready', 'submitted', 'accepted', 'returned') else None
        row['document_url'] = f"/api/packages/{package_id}/document" if row['download_url'] and row['manifest'].get('document_path') else None
        row['deletion_token'] = self._deletion_token(row, len(row['events']))
        return row

    @staticmethod
    def _deletion_token(row, receipt_count):
        return hashlib.sha256(dump([row['id'], row['state'], row['zip_hash'], receipt_count]).encode()).hexdigest()

    def delete(self, package_id, expected_token):
        with self._package_lock, self.db.tx() as con:
            row = con.execute('SELECT * FROM packages WHERE id=?', (package_id,)).fetchone()
            if not row:
                raise DomainError('报送记录不存在', 404)
            if row['state'] == 'deleted':
                return {'deleted': package_id, 'originals_retained': True}
            receipts = con.execute('SELECT COUNT(*) FROM package_events WHERE package_id=?', (package_id,)).fetchone()[0]
            if not expected_token or expected_token != self._deletion_token(row, receipts):
                raise DomainError('资料包状态或回执已变化，请重新打开删除确认框核对', 409)
            for job in con.execute("SELECT payload FROM jobs WHERE kind='export' AND state IN ('queued','running')"):
                payload = json.loads(job['payload'])
                groups = payload.get('groups', [])
                if (row['case_id'] in payload.get('package_ids', {}) or payload.get('parent_id') == package_id
                        or any(group.get('case', {}).get('id') == row['case_id'] for group in groups)
                        or not groups and not payload.get('package_ids')):
                    raise DomainError('本案资料包正在生成，请等待完成后再删除', 409)
            # Keep the immutable snapshot, receipt references and version reservation for audit/backup.
            con.execute("UPDATE packages SET state='deleted' WHERE id=?", (package_id,))
            audit(con, 'package', package_id, 'delete', {'state': row['state'], 'version': row['version'], 'case_id': row['case_id']},
                  {'state': 'deleted', 'originals_retained': True, 'receipts_retained': receipts})
        return {'deleted': package_id, 'originals_retained': True}

    def preview(self, photo_ids, template_id="general"):
        ids = list(dict.fromkeys(photo_ids))
        if not ids or len(ids) > 500:
            raise DomainError("每次请选择1至500个文件")
        rows = self.domain.photos(ids=ids, page_size=500)["items"]
        if len(rows) != len(ids):
            raise DomainError("部分照片不存在，请刷新列表")
        if template_id == PICC_FORMAT:
            case_ids = {row['case_id'] for row in rows}
            if len(case_ids) != 1 or None in case_ids:
                raise DomainError('人保记录表每次处理一个案件，请选择同一案件的资料')
            rows = self.picc.all_photos(next(iter(case_ids)))
            included = {p['id'] for p in self.picc.get(next(iter(case_ids)), rows)['photos']}
            rows = [p for p in rows if p['media_type'] != 'image' or p['id'] in included]
        excluded = [row for row in rows if row['media_type'] not in ('image', 'video')]
        rows = [row for row in rows if row['media_type'] in ('image', 'video')]
        if not rows:
            raise DomainError('所选资料没有照片或视频；压缩包不包含 PDF 或其他资料')
        if any(not p["case_id"] for p in rows):
            raise DomainError("所选照片含未归属案件的资料，请先确认案件")
        grouped = defaultdict(list)
        for row in rows:
            if template_id != PICC_FORMAT:
                photo_directory(row)
            grouped[row["case_id"]].append(row["id"])
        cases = []
        for case_id, values in grouped.items():
            case = self.domain.case(case_id)
            cases.append({'case': case, 'photo_ids': values, 'root_directory': case_folder(case),
                          'completeness': self.domain.completeness(case_id, template_id, values)})
        return {"photo_count": sum(row['media_type'] == 'image' for row in rows),
                "video_count": sum(row['media_type'] == 'video' for row in rows),
                "excluded_file_count": len(excluded), "cases": cases}

    def _content(self, case_id, photo_ids, template_id, parent_id):
        case = self.domain.case(case_id)
        template = next((t for t in self.domain.templates() if t["id"] == template_id), None)
        if not template:
            raise DomainError("资料要求模板不存在")
        if parent_id:
            parent = self.package(parent_id)
            if parent["case_id"] != case_id:
                raise DomainError("补件记录必须属于同一案件")
        is_picc = template_id == PICC_FORMAT
        photo_rows = self.picc.all_photos(case_id) if is_picc else [row for start in range(0, len(photo_ids), 500)
            for row in self.domain.photos(ids=photo_ids[start:start + 500], case_id=case_id, page_size=500)['items']]
        if not is_picc and len(photo_rows) != len(photo_ids):
            raise DomainError("照片所属案件已发生变化，请重新选材")
        photo_rows = [row for row in photo_rows if row['media_type'] == 'image' or not is_picc and row['media_type'] == 'video']
        if not photo_rows:
            raise DomainError('本案尚无照片，视频请使用原照片打包方案' if is_picc else '所选资料没有照片或视频；压缩包不包含 PDF 或其他资料')
        photo_ids = [row['id'] for row in photo_rows]
        record = self.picc.snapshot(case_id, photo_rows) if is_picc else None
        if is_picc:
            included = {p['id'] for p in record['photos']}
            photo_rows = [p for p in photo_rows if p['id'] in included]
            photo_ids = [p['id'] for p in photo_rows]
            root_directory = case_folder(case)
            paths = {row['id']: '_document_media/' + row['id'] + '.jpg' for row in photo_rows if row['media_type'] == 'image'}
            order = {photo['id']: index for index, photo in enumerate(record['photos'])}
            photo_rows.sort(key=lambda row: (order.get(row['id'], len(order)), row['id']))
        else:
            root_directory, paths = photo_paths(case, photo_rows)
        snapshots = []
        for row in photo_rows:
            source_rows = self.db.all("SELECT source_type,original_name,group_name,sender,message_time,received_at FROM photo_sources WHERE photo_id=? ORDER BY received_at", (row["id"],))
            original = row["original_name"] or (row["id"] + row["extension"])
            path = paths[row['id']]
            allowed = ("id", "blob_hash", "stage", "role", "photo_group", "part", "process", "watermark_time", "exif_time", "recognition_state", "confirmed_fields", "media_type", "mime_type", "size_bytes", "width", "height", "revision", "visit_id", "visit_date", "visit_key", "shared_checkin_visit_id", "shared_checkin_date", "shared_checkin_key")
            snapshots.append({**{k: row[k] for k in allowed}, "original_name": original, "package_path": path, "sources": source_rows})
        quality = self.domain.completeness(case_id, template_id, photo_ids)
        scope_visits = {r['scope_key'][6:] for r in quality['items'] if r['scope_key'].startswith('visit:')}
        visits = []
        for visit in self.domain.visits(case_id):
            selected = {**visit, 'photo_ids': [pid for pid in visit['photo_ids'] if pid in photo_ids],
                        'shared_checkin_ids': [pid for pid in visit['shared_checkin_ids'] if pid in photo_ids]}
            if selected['photo_ids'] or selected['shared_checkin_ids'] or visit['id'] in scope_visits:
                # Counts in a selected package must not imply excluded originals are included.
                selected.pop('counts', None)
                visits.append(selected)
        return {"case": case, "template": template, "photos": snapshots, "completeness": quality,
                "visits": visits, "parent_id": parent_id, "format": PICC_FORMAT if is_picc else PHOTO_PACKAGE_FORMAT,
                "root_directory": root_directory,
                **({'record': record, 'document_path': root_directory + '/查勘及施工记录表.docx'} if is_picc else {})}

    def _equivalent(self, content, con, *, include_building=False):
        fingerprint = content_fingerprint(content)
        rows = con.execute("SELECT * FROM packages WHERE case_id=? AND state IN ('ready','submitted','accepted','returned'" +
                           (",'building'" if include_building else "") + ") ORDER BY version DESC", (content['case']['id'],))
        for row in rows:
            manifest = json.loads(row['manifest'])
            if content_fingerprint(manifest) != fingerprint:
                continue
            if row['state'] != 'building':
                path = self.storage.resolve(row['zip_path'])
                if not path.is_file():
                    continue
                with path.open('rb') as stream:
                    if hashlib.file_digest(stream, 'sha256').hexdigest() != row['zip_hash']:
                        continue
                if manifest.get('document_path'):
                    document = self.storage.resolve('exports/' + row['id'] + '/' + manifest['document_path'])
                    if not document.is_file():
                        continue
                    with document.open('rb') as stream:
                        if hashlib.file_digest(stream, 'sha256').hexdigest() != manifest.get('document_export', {}).get('sha256'):
                            continue
            return dict(row)
        return None

    @staticmethod
    def _export_result(package, reused=False):
        return {"id": package['id'], "case_id": package['case_id'], "version": package['version'],
                "download_url": package['download_url'], "document_url": package.get('document_url'), "reused": reused}

    def enqueue_export(self, jobs, photo_ids, template_id='general', parent_id=None, title=''):
        preview = self.preview(photo_ids, template_id)
        return self._enqueue_groups(jobs, preview['cases'], template_id, parent_id, title)

    def enqueue_case_export(self, jobs, case_id, template_id, title=''):
        # Internal all-case preparation is not limited to one UI selection page.
        ids = [p['id'] for p in self.picc.all_photos(case_id) if p['media_type'] in ('image', 'video')]
        return self._enqueue_groups(jobs, [{'case': {'id': case_id}, 'photo_ids': ids}], template_id, None, title)

    def _enqueue_groups(self, jobs, groups, template_id, parent_id, title):
        with self._package_lock:
            with self.db.read_snapshot() as con:
                if parent_id and len(groups) != 1:
                    raise DomainError('补件请每次选择一个案件')
                contents = [self._content(group['case']['id'], group['photo_ids'], template_id, parent_id) for group in groups]
                fingerprint = hashlib.sha256(dump(sorted((item['case']['id'], content_fingerprint(item)) for item in contents)).encode()).hexdigest()
                for job in con.execute("SELECT id,payload FROM jobs WHERE kind='export' AND state IN ('queued','running') ORDER BY created_at DESC"):
                    if json.loads(job['payload']).get('request_fingerprint') == fingerprint:
                        return jobs.get(job['id'])
                matches = [self._equivalent(content, con) for content in contents]
                results = [self._export_result(self.package(row['id']), True) for row in matches if row]
            payload = {"photo_ids": [pid for g in groups for pid in g['photo_ids']], "template_id": template_id, "parent_id": parent_id, "title": title,
                       "package_ids": {g['case']['id']: uid() for g in groups}, "request_fingerprint": fingerprint,
                       "groups": [{"case": {"id": g['case']['id']}, "photo_ids": g['photo_ids']} for g in groups]}
            if len(results) == len(contents):
                return jobs.enqueue('export', payload, result={"state": "completed", "packages": results, "reused_count": len(results)})
            return jobs.enqueue('export', payload)

    def _snapshot(self, package_id, case_id, photo_ids, template_id, parent_id, title):
        existing = self.db.one("SELECT * FROM packages WHERE id=?", (package_id,))
        if existing:
            if existing['state'] == 'deleted':
                raise DomainError('资料包已删除，请重新选材生成', 409)
            return json.loads(existing["manifest"])
        with self.db.read_snapshot():
            content = self._content(case_id, photo_ids, template_id, parent_id)
        with self.db.tx() as con:
            if parent_id:
                parent = con.execute("SELECT state FROM packages WHERE id=?", (parent_id,)).fetchone()
                if not parent or parent['state'] == 'deleted':
                    raise DomainError('原报送记录已删除，请重新选材生成', 409)
            equivalent = self._equivalent(content, con, include_building=True)
            if equivalent:
                return json.loads(equivalent['manifest'])
            version = con.execute("SELECT COALESCE(MAX(version),0)+1 FROM packages WHERE case_id=?", (case_id,)).fetchone()[0]
            manifest = {**content, "schema": 1, "package_id": package_id, "version": version, "created_at": now(),
                        "evidence_note": "按本次选材形成的不可变快照；完工照片不等同于验收通过；生成报送包不等同于提交。"}
            zip_path = f"exports/{package_id}/资料报送包.zip"
            con.execute("INSERT INTO packages VALUES(?,?,?,?,?,?,?,?,?,?,?)", (package_id, case_id, template_id, parent_id, version, title or content['case']["title"], "building", dump(manifest), zip_path, "", now()))
            audit(con, "package", package_id, "snapshot", after={"case_id": case_id, "version": version, "photo_ids": photo_ids})
        return manifest

    def build(self, package_id, case_id, photo_ids, template_id="general", parent_id=None, title=""):
        with self._package_lock:
            return self._build(package_id, case_id, photo_ids, template_id, parent_id, title)

    def _build(self, package_id, case_id, photo_ids, template_id, parent_id, title):
        old = self.db.one("SELECT state FROM packages WHERE id=?", (package_id,))
        if old and old["state"] in ("ready", "submitted", "accepted", "returned"):
            return self.package(package_id)
        manifest = self._snapshot(package_id, case_id, photo_ids, template_id, parent_id, title)
        reused = manifest['package_id'] != package_id
        if reused:
            package = self.package(manifest['package_id'])
            if package['state'] != 'building':
                return {**package, 'reused': True}
            # Finish an interrupted, identical snapshot instead of completing a job with no ZIP.
            package_id = manifest['package_id']
        folder = self.storage.resolve("exports/" + package_id)
        folder.mkdir(parents=True, exist_ok=True)
        try:
            media_only = manifest.get('format') == PHOTO_PACKAGE_FORMAT
            is_picc = manifest.get('format') == PICC_FORMAT
            for photo in manifest["photos"]:
                blob = self.db.one("SELECT * FROM blobs WHERE hash=?", (photo["blob_hash"],))
                if not blob or not self.storage.verify(blob):
                    raise DomainError("原文件缺失或哈希不一致：" + photo["original_name"])
                destination = folder / photo["package_path"]
                destination.parent.mkdir(parents=True, exist_ok=True)
                if media_only or is_picc:
                    photo['export'] = write_media(self.storage.resolve(blob['stored_path']), destination, photo, force_jpeg=is_picc)
                elif not destination.is_file() or hashlib.sha256(destination.read_bytes()).hexdigest() != photo["blob_hash"]:
                    shutil.copyfile(self.storage.resolve(blob["stored_path"]), destination)
            if is_picc:
                manifest['document_export'] = write_record(folder / manifest['document_path'], manifest['record'], manifest['photos'], folder)
            (folder / "snapshot.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
            if is_picc:
                # A PICC delivery is the Word document itself. Videos use the ordinary media export.
                document_path = 'exports/' + package_id + '/' + manifest['document_path']
                with self.db.tx() as con:
                    con.execute("UPDATE packages SET state='ready',zip_path=?,zip_hash=?,manifest=? WHERE id=? AND state<>'deleted'",
                                (document_path, manifest['document_export']['sha256'], dump(manifest), package_id))
                return {**self.package(package_id), 'reused': reused}
            if not media_only and not is_picc:
                # An interrupted historical job must keep its original immutable export format.
                spreadsheet(folder / "照片资料索引.xlsx", manifest)
                word_report(folder / "案件照片资料报告.docx", manifest, folder)
                html_report(folder / "照片浏览.html", manifest)
                missing = [f"{STAGES[r['stage']]} / {ROLES[r['role']]} / {r['scope']}：要求{r['minimum']}，本次{r['count']}" for r in manifest["completeness"]["items"] if r["missing"]]
                (folder / "缺项说明.txt").write_text("本次选材核对\n" + ("\n".join(missing) if missing else "本次选材满足当前模板数量要求") + "\n\n" + manifest["evidence_note"], encoding="utf-8")
            target = folder / "资料报送包.zip"
            temp = folder / "资料报送包.zip.tmp"
            file_paths = [folder / p["package_path"] for p in manifest["photos"]]
            if not media_only:
                file_paths += [folder / name for name in ("snapshot.json", "照片资料索引.xlsx", "案件照片资料报告.docx", "照片浏览.html", "缺项说明.txt")]
            with zipfile.ZipFile(temp, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                if media_only:
                    root = manifest['root_directory']
                    for directory in ('', *PHOTO_DIRECTORIES):
                        archive.writestr(root + ('/' + directory if directory else '') + '/', b'')
                for path in file_paths:
                    archive.write(path, path.relative_to(folder).as_posix())
            os.replace(temp, target)
            digest = hashlib.sha256(target.read_bytes()).hexdigest()
            with self.db.tx() as con:
                con.execute("UPDATE packages SET state='ready',zip_hash=?,manifest=? WHERE id=? AND state<>'deleted'", (digest, dump(manifest), package_id))
            return {**self.package(package_id), 'reused': reused}
        except Exception:
            with self.db.tx() as con:
                con.execute("UPDATE packages SET state='failed' WHERE id=? AND state='building'", (package_id,))
            raise

    def process(self, payload, progress):
        groups = payload.get("groups")
        if groups is None:
            groups = self.preview(payload["photo_ids"], payload.get("template_id", "general"))["cases"]
        mapping = payload.get("package_ids", {})
        results = []
        for index, group in enumerate(groups):
            case_id = group["case"]["id"]
            package_id = mapping.get(case_id) or uid()
            result = self.build(package_id, case_id, group["photo_ids"], payload.get("template_id", "general"), payload.get("parent_id"), payload.get("title", ""))
            results.append(self._export_result(result, result.get('reused', False)))
            progress(index + 1, len(groups))
        return {"state": "completed", "packages": results, "reused_count": sum(result['reused'] for result in results)}

    def add_event(self, package_id, data, attachment_hash=None):
        package = self.package(package_id)
        if package["state"] in ("building", "failed"):
            raise DomainError("资料包尚未生成完成")
        state = data.get("state")
        if state not in ("submitted", "accepted", "returned"):
            raise DomainError("请选择已提交、已接收或退回补件")
        event_at = data.get("event_at") or now()
        from .domain import clean_date
        event_at = clean_date(event_at, "记录")
        with self.db.tx() as con:
            current = con.execute('SELECT state FROM packages WHERE id=?', (package_id,)).fetchone()
            if not current or current['state'] == 'deleted':
                raise DomainError('报送记录已删除', 404)
            con.execute("INSERT INTO package_events VALUES(?,?,?,?,?,?,?)", (uid(), package_id, state, event_at, str(data.get("note", ""))[:10000], attachment_hash, now()))
            con.execute("UPDATE packages SET state=? WHERE id=?", (state, package_id))
            audit(con, "package", package_id, "receipt", after={"state": state, "event_at": event_at, "note": data.get("note", ""), "attachment_hash": attachment_hash})
        return self.package(package_id)

    def backup(self, payload, progress):
        backup_id = payload.get("backup_id") or uid()
        root = self.storage.settings.data_dir
        snapshot_db = root / "staging" / (backup_id + ".sqlite3")
        self.db.backup(snapshot_db)
        with closing(sqlite3.connect(snapshot_db)) as con:
            files = [(r[0], r[1]) for r in con.execute("SELECT stored_path,hash FROM blobs")]
            files += [(r[0], None) for r in con.execute("SELECT thumbnail_path FROM blobs WHERE thumbnail_path IS NOT NULL")]
            package_ids = [r[0] for r in con.execute("SELECT id FROM packages")]
            package_ids += [r[0] for r in con.execute("SELECT id FROM platform_packages")]
        for package_id in package_ids:
            folder = root / "exports" / package_id
            if folder.is_dir():
                files += [(p.relative_to(root).as_posix(), None) for p in folder.rglob("*") if p.is_file() and not p.name.endswith(".tmp")]
        target = root / "backups" / ("idi-backup-" + backup_id + ".zip")
        temp = target.with_suffix(".zip.tmp")
        metadata = {"schema": 1, "created_at": now(), "files": {}}
        try:
            with zipfile.ZipFile(temp, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
                archive.write(snapshot_db, "archive.sqlite3")
                metadata["files"]["archive.sqlite3"] = hashlib.sha256(snapshot_db.read_bytes()).hexdigest()
                for index, (relative, expected_hash) in enumerate(files):
                    path = self.storage.resolve(relative)
                    if not path.is_file():
                        if expected_hash:
                            raise DomainError("备份失败：原件缺失 " + relative)
                        continue
                    digest = hashlib.sha256(path.read_bytes()).hexdigest()
                    if expected_hash and digest != expected_hash:
                        raise DomainError("备份失败：原件哈希不一致 " + relative)
                    archive.write(path, relative)
                    metadata["files"][relative] = digest
                    progress(index + 1, len(files))
                archive.writestr("backup-metadata.json", json.dumps(metadata, ensure_ascii=False, indent=2))
            os.replace(temp, target)
        finally:
            snapshot_db.unlink(missing_ok=True)
            temp.unlink(missing_ok=True)
        return {"state": "completed", "backup_id": backup_id, "filename": target.name, "download_url": f"/api/backups/{backup_id}/download", "file_count": len(metadata["files"])}

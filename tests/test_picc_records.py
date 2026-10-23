import hashlib
import io
import zipfile
import sqlite3
import json
from datetime import datetime, timedelta

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image
import pytest

from idi.db import uid, now, dump
from idi.picc_records import PICC_FORMAT, TEMPLATE_PATH, normalize_time


@pytest.fixture
def case(case, services):
    return services['domain'].save_case({'platform_report_code': 'ZLBH202610132005010577', 'insurer_template_id': PICC_FORMAT}, case['id'])


def plan(client, case, **changes):
    url = f"/api/cases/{case['id']}/picc-record"
    draft = client.get(url).json()
    for field, value in changes.items():
        if field == 'photo_overrides':
            draft[field] = value
        else:
            draft['fields'][field] = value
    saved = client.put(url, json=draft)
    assert saved.status_code == 200, saved.text
    return saved.json()


def build(client, services, ids):
    queued = client.post('/api/packages', json={'photo_ids': ids, 'template_id': PICC_FORMAT})
    assert queued.status_code == 200, queued.text
    while services['jobs'].run_one():
        pass
    job = services['jobs'].get(queued.json()['id'])
    assert job['state'] == 'completed', job
    return services['reporter'].package(job['result']['packages'][0]['id'])


def test_record_fills_headers_all_photos_order_and_only_downloads_word(client, services, case, photo):
    late = photo('red', name='晚.png', watermark_time='2026-10-23T10:00:00', exif_time='2026-10-23T07:00:00')
    early = photo('blue', name='早.png', part='不能填入备注的部位', process='不能填入备注的施工内容')
    # EXIF is read from original files, not import metadata; seed the extracted values.
    with services['db'].tx() as con:
        con.execute('UPDATE photos SET exif_time=? WHERE id=?', ('2026-10-23T07:00:00', late['id']))
        con.execute('UPDATE photos SET exif_time=?,created_at=? WHERE id=?', ('2026-10-23T06:00:00+08:00', '2026-10-23T01:00:00Z', early['id']))
    uploaded = photo('green', name='上传.png')
    with services['db'].tx() as con:
        con.execute('UPDATE photos SET created_at=? WHERE id=?', ('2026-10-23T00:30:00Z', uploaded['id']))
    video = services['importer'].add_file(b'video-original' * 40000, '记录视频.mp4', case_id=case['id'])['photo_id']
    services['importer'].add_file(b'%PDF-1.4\nfixture', '资料.pdf', case_id=case['id'])
    draft = plan(client, case, survey_date='1999-01-01', survey_notes='旧查勘说明', record_number='3', photo_overrides={uploaded['id']: {'sort_time': '2026-10-23T12:30', 'remarks': '旧备注'}})
    assert [p['id'] for p in draft['photos']] == [uploaded['id'], early['id'], late['id']]
    assert [p['time_source'] for p in draft['photos']] == ['upload', 'upload', 'watermark']
    assert draft['fields']['survey_date'] == '2026-10-23'
    assert not draft['photo_overrides'] and all(not p['remarks'] for p in draft['photos'])
    preview = services['reporter'].preview([late['id']], PICC_FORMAT)
    assert (preview['photo_count'], preview['video_count'], preview['excluded_file_count']) == (3, 1, 1)
    before = {p['id']: services['domain'].photo(p['id']) for p in (late, early, uploaded)}
    package = build(client, services, [late['id']])
    manifest = package['manifest']
    assert [p['id'] for p in manifest['photos'][:3]] == [uploaded['id'], early['id'], late['id']]
    result = client.get(package['document_url'])
    assert result.status_code == 200 and '.docx' in result.headers['content-disposition']
    document = Document(io.BytesIO(result.content))
    table = document.tables[0]
    assert len(table.rows) == 7 and len(document.inline_shapes) == 3
    assert table.rows[0].cells[0].text == '查勘记录表'
    assert '从化铂悦府' in table.rows[1].cells[0].text
    assert '2026-10-23' in table.rows[2].cells[1].text
    assert table.rows[3].cells[1].text == '备注'
    assert all(row.cells[1].text == '' for row in table.rows[4:])
    alternate = client.get(package['download_url'])
    assert alternate.content == result.content and '.docx' in alternate.headers['content-disposition']
    assert len(manifest['photos']) == 3
    assert not any(p.suffix in ('.mp4', '.zip') for p in services['storage'].resolve('exports/' + package['id']).rglob('*'))
    video_photo = services['domain'].photo(video)
    assert services['storage'].resolve(services['db'].one('SELECT stored_path FROM blobs WHERE hash=?', (video_photo['blob_hash'],))['stored_path']).read_bytes() == b'video-original' * 40000
    for p in (late, early, uploaded):
        assert services['domain'].photo(p['id']) == before[p['id']]
    assert services['domain'].photos(case_id=case['id'])['total'] == 5


def test_upload_time_fills_missing_watermark_and_date_without_manual_steps(client, services, case, photo):
    p = photo()
    with services['db'].tx() as con:
        con.execute('UPDATE photos SET created_at=? WHERE id=?', ('2026-10-22T18:01:02.321Z', p['id']))
    url = f"/api/cases/{case['id']}/picc-record"
    draft = client.get(url).json()
    assert draft['photos'][0]['sort_time'] == '2026-10-23T02:01:02.321000' and draft['missing_time_count'] == 0
    assert draft['photos'][0]['time_source'] == 'upload'
    assert draft['fields']['survey_date'] == '2026-10-23' and draft['fields']['survey_notes'] == ''
    assert draft['blockers'] == [] and draft['revision'] == 0
    assert build(client, services, [p['id']])['state'] == 'ready'


def test_draft_revision_validation_and_ignored_legacy_overrides(client, services, case, photo, project):
    p = photo(watermark_time='2026-10-23T10:00:00')
    draft = plan(client, case)
    url = f"/api/cases/{case['id']}/picc-record"
    stale = {**draft, 'revision': draft['revision'] - 1}
    assert client.put(url, json=stale).status_code == 409
    draft['photo_overrides'] = []
    assert client.put(url, json=draft).status_code == 400
    draft['photo_overrides'] = {}
    draft['fields']['project_name'] = 'x' * 121
    assert client.put(url, json=draft).status_code == 400
    draft['fields']['project_name'] = '原项目'
    other = services['domain'].save_case({'project_id': project['id'], 'title': '其他案件'})
    services['domain'].update_photo(p['id'], {'case_id': other['id']})
    draft['photo_overrides'] = {p['id']: {'remarks': '错误案件'}}
    saved = client.put(url, json=draft)
    assert saved.status_code == 200 and saved.json()['photo_overrides'] == {}
    assert client.get(url).json()['photos'] == []


def test_content_reuse_changed_header_new_version_and_deleted_document(client, services, case, photo):
    p = photo(watermark_time='2026-10-23T10:00:00')
    plan(client, case)
    first = build(client, services, [p['id']])
    original = client.get(first['document_url']).content
    plan(client, case)
    assert build(client, services, [p['id']])['id'] == first['id']
    plan(client, case, survey_content='核对后的查勘内容')
    second = build(client, services, [p['id']])
    assert second['version'] == 2 and client.get(first['document_url']).content == original
    assert client.request('DELETE', '/api/packages/' + second['id'], json={'expected_token': second['deletion_token']}).status_code == 200
    assert client.get(second['document_url']).status_code == 404
    assert services['domain'].photo(p['id'])


def test_date_sources_are_automatic_and_refresh_from_case_information(client, services, case, photo):
    construction = photo('red', watermark_time='2026-10-01T09:00:00', stage='work', role='construction')
    surveyed = photo('blue', watermark_time='2026-10-18T00:10:00Z', stage='survey', role='problem')
    url = f"/api/cases/{case['id']}/picc-record"
    draft = client.get(url).json()
    assert draft['fields']['survey_date'] == '2026-10-18'
    assert draft['survey_date_source'] == '最早查勘照片的水印时间'
    # Existing manual values from the first version must not override automatic rules.
    with services['db'].tx() as con:
        con.execute('INSERT INTO picc_record_drafts VALUES(?,?,?,?,?,?)',
                    (case['id'], dump({**draft['fields'], 'survey_date': '1999-01-01', 'survey_notes': '旧说明'}),
                     dump({surveyed['id']: {'sort_time': '1999-01-01T01:00:00', 'remarks': '旧备注'}}), 1, now(), now()))
    legacy = services['db'].one('SELECT * FROM picc_record_drafts WHERE case_id=?', (case['id'],))
    refreshed = client.get(url).json()
    assert refreshed['fields']['survey_date'] == '2026-10-18' and refreshed['fields']['survey_notes'] == ''
    assert refreshed['photos'][0]['id'] == construction['id']
    assert all(p['remarks'] == '' for p in refreshed['photos'])
    assert services['db'].one('SELECT * FROM picc_record_drafts WHERE case_id=?', (case['id'],)) == legacy
    services['domain'].save_visit(case['id'], {'stage': 'survey', 'visit_date': '2026-10-17'})
    assert client.get(url).json()['fields']['survey_date'] == '2026-10-17'
    with services['db'].tx() as con:
        con.execute('INSERT INTO platform_links(case_id,report_receiving_id,snapshot,updated_at) VALUES(?,?,?,?)',
                    (case['id'], 'survey-report', '{}', now()))
        con.execute('INSERT INTO platform_drafts(case_id,stage,values_json,target_report_id,updated_at) VALUES(?,?,?,?,?)',
                    (case['id'], 'survey', dump({'siteSurveyDate': '2026-10-16 09:00:00'}), 'survey-report', now()))
    assert client.get(url).json()['fields']['survey_date'] == '2026-10-16'
    with services['db'].tx() as con:
        con.execute('UPDATE platform_links SET snapshot=? WHERE case_id=?',
                    (dump({'survey': {'siteSurveyDate': '2026-10-14T18:30:00Z'}}), case['id']))
    plan = client.get(url).json()
    assert plan['fields']['survey_date'] == '2026-10-15' and plan['survey_date_source'] == '官网查勘信息'
    package = build(client, services, [surveyed['id']])
    doc = Document(io.BytesIO(client.get(package['document_url']).content))
    assert doc.tables[0].rows[2].cells[1].text == '查勘日期：2026-10-15'
    assert all(row.cells[1].text == '' for row in doc.tables[0].rows[4:])


def test_invalid_watermark_uses_upload_and_preserves_subsecond_order(client, services, case, photo):
    early = photo('red', name='Z.png')
    late = photo('blue', name='A.png')
    with services['db'].tx() as con:
        for p, millis in ((early, '100'), (late, '200')):
            con.execute('UPDATE photos SET watermark_time=?,created_at=?,exif_time=? WHERE id=?',
                        ('invalid', f'2026-10-23T01:00:00.{millis}Z', '1999-01-01T00:00:00', p['id']))
    result = client.get(f"/api/cases/{case['id']}/picc-record").json()
    assert [p['id'] for p in result['photos']] == [early['id'], late['id']]
    assert all(p['time_source'] == 'upload' for p in result['photos'])
    assert result['fields']['survey_date'] == '2026-10-23' and not result['blockers']


def test_case_information_changes_create_new_record_version(client, services, case, photo):
    p = photo(watermark_time='2026-10-23T10:00:00')
    first = build(client, services, [p['id']])
    original = client.get(first['document_url']).content
    services['domain'].save_visit(case['id'], {'stage': 'survey', 'visit_date': '2026-10-22'})
    second = build(client, services, [p['id']])
    assert second['version'] == 2 and client.get(first['document_url']).content == original
    assert second['manifest']['record']['fields']['survey_date'] == '2026-10-22'
    assert second['manifest']['record']['rules_version'] == 3


@pytest.mark.parametrize('damage', ['missing', 'corrupt'])
def test_document_damage_detected_and_regenerated(client, services, case, photo, damage):
    p = photo(watermark_time='2026-10-23T10:00:00')
    plan(client, case)
    first = build(client, services, [p['id']])
    path = services['storage'].resolve('exports/' + first['id'] + '/' + first['manifest']['document_path'])
    if damage == 'missing': path.unlink()
    else: path.write_bytes(b'damaged')
    assert client.get(first['document_url']).status_code in (404, 409)
    assert build(client, services, [p['id']])['version'] == 2


def test_template_is_sanitized_preserves_geometry_styles_and_image_cap(client, services, case):
    pixels = Image.effect_noise((1300, 900), 90).convert('RGB')
    stream = io.BytesIO()
    pixels.save(stream, 'PNG')
    raw = stream.getvalue()
    assert len(raw) > 400000
    p = services['importer'].add_file(raw, '大照片.png', case_id=case['id'], metadata={'watermark_time': '2026-10-23T10:00:00'})['photo_id']
    plan(client, case)
    package = build(client, services, [p])
    result = client.get(package['document_url']).content
    with zipfile.ZipFile(TEMPLATE_PATH) as template, zipfile.ZipFile(io.BytesIO(result)) as archive:
        assert not any(n.startswith('word/media/') for n in template.namelist())
        assert '振业颐隽' not in template.read('word/document.xml').decode()
        assert '非人为因素' not in archive.read('word/document.xml').decode()
        for name in ['word/styles.xml', 'word/theme/theme1.xml', 'word/fontTable.xml', 'word/settings.xml', 'word/webSettings.xml']:
            assert archive.read(name) == template.read(name)
        images = [archive.read(n) for n in archive.namelist() if n.startswith('word/media/')]
        assert len(images) == 1 and len(images[0]) <= 400000
        assert Image.open(io.BytesIO(images[0])).format == 'JPEG'
    blob = services['db'].one('SELECT * FROM blobs WHERE hash=?', (hashlib.sha256(raw).hexdigest(),))
    assert services['storage'].resolve(blob['stored_path']).read_bytes() == raw
    source = Document(TEMPLATE_PATH).sections[0]
    target = Document(io.BytesIO(result)).sections[0]
    assert (source.page_width, source.page_height, source.left_margin, source.top_margin) == (target.page_width, target.page_height, target.left_margin, target.top_margin)


def test_more_than_500_photos_all_included_in_plan_and_snapshot(services, case, photo):
    p = photo(watermark_time='2026-10-23T10:00:00', stage='survey', role='problem')
    with services['db'].tx() as con:
        row = dict(con.execute('SELECT * FROM photos WHERE id=?', (p['id'],)).fetchone())
        blob = dict(con.execute('SELECT * FROM blobs WHERE hash=?', (row['blob_hash'],)).fetchone())
        columns = ','.join(row)
        for index in range(500):
            digest = hashlib.sha256(str(index).encode()).hexdigest()
            blob_clone = {**blob, 'hash': digest, 'stored_path': f'isolated-fixture/{index}.png'}
            con.execute(f"INSERT INTO blobs({','.join(blob_clone)}) VALUES({','.join('?' for _ in blob_clone)})", tuple(blob_clone.values()))
            clone = {**row, 'id': uid(), 'blob_hash': digest, 'watermark_time': (datetime(2026, 10, 24) + timedelta(minutes=index)).isoformat()}
            con.execute(f"INSERT INTO photos({columns}) VALUES({','.join('?' for _ in clone)})", tuple(clone.values()))
    record = services['reporter'].picc.get(case['id'])
    record['fields']['survey_date'] = '2026-10-23'
    services['reporter'].picc.save(case['id'], record)
    preview = services['reporter'].preview([p['id']], PICC_FORMAT)
    assert preview['photo_count'] == 501
    content = services['reporter']._content(case['id'], [p['id']], PICC_FORMAT, None)
    assert len(content['record']['photos']) == len(content['photos']) == 501
    ordinary = services['reporter']._content(case['id'], [p['id'] for p in record['photos']], 'insurer-v2', None)
    assert len(ordinary['photos']) == 501


def test_multi_case_and_document_authorization(client, app, services, case, photo, project):
    p = photo(watermark_time='2026-10-23T10:00:00')
    other = services['domain'].save_case({'project_id': project['id'], 'title': '另案'})
    q = services['importer'].add_file(b'video', 'other.mp4', case_id=other['id'])['photo_id']
    response = client.post('/api/packages/preview', json={'photo_ids': [p['id'], q], 'template_id': PICC_FORMAT})
    assert response.status_code == 400 and '一个案件' in response.text
    plan(client, case)
    package = build(client, services, [p['id']])
    with TestClient(app) as anonymous:
        assert anonymous.get(package['document_url']).status_code == 401
        assert anonymous.get(f"/api/cases/{case['id']}/picc-record").status_code == 401


def test_draft_project_scope_and_document_backup(client, app, services, case, photo, tmp_path):
    p = photo(watermark_time='2026-10-23T10:00:00')
    saved = plan(client, case, survey_content='备份验收')
    package = build(client, services, [p['id']])
    user = services['access'].save_user({'username': 'record-worker', 'display_name': '项目人员', 'password': 'person-test-password', 'system_role': 'user'})
    with TestClient(app, headers={'X-IDI-Request': '1'}) as worker:
        assert worker.post('/api/auth/login', json={'username': 'record-worker', 'password': 'person-test-password'}).status_code == 200
        url = f"/api/cases/{case['id']}/picc-record"
        assert worker.get(url).status_code == 403
        assert worker.put(url, json=saved).status_code == 403
        services['access'].save_members(case['project_id'], {'members': [{'user_id': user['id'], 'role': 'member'}]})
        assert worker.get(url).status_code == 200
        assert worker.put(url, json=saved).status_code == 403
        assert worker.get(package['document_url']).status_code == 403
    backup = services['reporter'].backup({}, lambda *_: None)
    with zipfile.ZipFile(services['storage'].resolve('backups/' + backup['filename'])) as archive:
        doc_path = 'exports/' + package['id'] + '/' + package['manifest']['document_path']
        assert archive.read(doc_path) == client.get(package['document_url']).content
        target = tmp_path / 'backup-db.sqlite3'
        target.write_bytes(archive.read('archive.sqlite3'))
    with sqlite3.connect(target) as con:
        assert '备份验收' in con.execute('SELECT fields FROM picc_record_drafts WHERE case_id=?', (case['id'],)).fetchone()[0]


@pytest.mark.parametrize(('value', 'expected'), [('2026-10-23T01:00:00Z', '2026-10-23T09:00:00'), ('2026-10-23', None), ('bad', None), ('2026-10-23 09:00', '2026-10-23T09:00:00')])
def test_timestamp_normalization(value, expected):
    assert normalize_time(value) == expected


def test_duplicate_copies_removed_as_one_photo_and_restored_without_original_changes(client, services, case, photo):
    p = photo('red', name='同名照片.png', stage='survey', role='building_checkin', watermark_time='2026-10-23T08:00:00')
    q = photo('blue', name='同名照片.png', watermark_time='2026-10-23T09:00:00')
    endpoint = f"/api/cases/{case['id']}/picc-record"
    before = {pid: services['domain'].photo(pid) for pid in (p['id'], q['id'])}
    first = build(client, services, [p['id']])
    original = client.get(first['document_url']).content
    draft = client.get(endpoint).json()
    removed = client.put(endpoint, json={**draft, 'photo_action': {'id': p['id'], 'excluded': True}})
    assert removed.status_code == 200
    assert [item['id'] for item in removed.json()['photos']] == [q['id']]
    # A new work copy of an already excluded image must not reappear in the record.
    copied = services['domain'].reuse_checkin(p['id'], {'case_id': case['id'], 'revision': p['revision']})
    draft = client.get(endpoint).json()
    assert len(draft['photos']) == len(draft['excluded_photos']) == draft['duplicate_count'] == 1
    assert draft['excluded_photos'][0]['id'] == p['id']
    assert services['reporter'].preview([p['id']], PICC_FORMAT)['photo_count'] == 1
    second = build(client, services, [p['id']])
    assert len(second['manifest']['photos']) == 1 and second['version'] == 2
    assert len(Document(io.BytesIO(client.get(second['document_url']).content)).inline_shapes) == 1
    assert client.get(first['document_url']).content == original
    # Saving only header fields, including by an old client, preserves exclusions.
    again = client.put(endpoint, json={'revision': draft['revision'], 'fields': draft['fields']}).json()
    assert len(again['excluded_photos']) == 1
    restored = client.put(endpoint, json={**again, 'photo_action': {'id': p['id'], 'excluded': False}}).json()
    assert len(restored['photos']) == 2 and restored['excluded_photos'] == []
    assert restored['duplicate_count'] == 1  # Different image with the same name stays.
    third = build(client, services, [q['id']])
    assert len(third['manifest']['photos']) == 2
    assert services['domain'].photos(case_id=case['id'])['total'] == 3
    for pid, old in before.items():
        current = services['domain'].photo(pid)
        assert (current['blob_hash'], current['revision'], current['stage'], current['role']) == (old['blob_hash'], old['revision'], old['stage'], old['role'])


def test_photo_removal_validates_case_revision_and_all_excluded(client, services, case, photo, project):
    p = photo()
    other = services['domain'].save_case({'project_id': project['id'], 'title': '另案'})
    endpoint = f"/api/cases/{case['id']}/picc-record"
    draft = client.get(endpoint).json()
    assert client.put(endpoint, json={**draft, 'photo_action': {'id': p['id'], 'excluded': 'yes'}}).status_code == 400
    removed = client.put(endpoint, json={**draft, 'photo_action': {'id': p['id'], 'excluded': True}}).json()
    assert removed['photos'] == [] and '全部移除' in removed['blockers'][0]
    assert client.put(endpoint, json={**draft, 'photo_action': {'id': p['id'], 'excluded': False}}).status_code == 409
    other_draft = client.get(f"/api/cases/{other['id']}/picc-record").json()
    assert client.put(f"/api/cases/{other['id']}/picc-record", json={**other_draft, 'photo_action': {'id': p['id'], 'excluded': True}}).status_code == 409

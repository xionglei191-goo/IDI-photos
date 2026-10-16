import hashlib
import io
import json
import zipfile
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from idi.db import Database, dump, uid
from idi.domain import DomainError, GROUP_FIELDS


def entry(**changes):
    return {'stage': 'work', 'visit_date': '2026-10-16', 'visit_key': '第一次进场',
            'started_at': '2026-10-16T09:29', 'finished_at': '2026-10-16T09:47',
            'crew_number': 1, 'same_day_repair': True, 'notes': '本次维修记录', **changes}


def materials(photo):
    return [photo(color, name, stage=stage, role=role, watermark_time='2026-10-16 ' + time)
            for color, name, stage, role, time in [
                ('red', '楼栋.png', 'survey', 'building_checkin', '09:25'),
                ('blue', '入户.png', 'survey', 'door_checkin', '09:29'),
                ('green', '施工.png', 'work', 'construction', '09:35'),
                ('yellow', '完工.png', 'acceptance', 'finished', '09:47')]]


def record(domain, case, photos, **changes):
    return domain.save_visit(case['id'], entry(photo_ids=[p['id'] for p in photos[2:]],
                                             shared_checkin_ids=[p['id'] for p in photos[:2]], **changes))


def test_shared_originals_count_in_two_uses_but_remain_four_originals(services, case, photo, client):
    domain = services['domain']
    photos = materials(photo)
    originals = [services['storage'].resolve(services['db'].one('SELECT stored_path FROM blobs WHERE hash=?', (p['blob_hash'],))['stored_path']).read_bytes() for p in photos]
    before = domain.completeness(case['id'])
    assert any(r['missing'] for r in before['items'] if r['stage'] == 'work' and 'checkin' in r['role'])
    visit = record(domain, case, photos)
    assert visit['counts'] == {'building_checkin': 1, 'door_checkin': 1, 'construction': 1, 'completion': 1, 'shared_checkin': 2}
    assert domain.photos(case_id=case['id'])['total'] == 4
    assert domain.photos(case_id=case['id'], photo_group='survey_checkin')['total'] == 2
    assert domain.photos(case_id=case['id'], photo_group='work_checkin')['total'] == 2
    assert domain.photos(case_id=case['id'], stage='work')['total'] == 3
    response = client.get('/api/photos', params={'case_id': case['id'], 'visit_id': visit['id']}).json()
    assert {p['id'] for p in response['items']} == {p['id'] for p in photos}
    for p, data in zip(photos, originals):
        current = domain.photo(p['id'])
        assert current['blob_hash'] == hashlib.sha256(data).hexdigest()
        assert (current['stage'], current['role'], current['photo_group']) == (p['stage'], p['role'], p['photo_group'])
    for p in photos[:2]:
        assert domain.photo(p['id'])['visit_id'] is None
        assert domain.photo(p['id'])['shared_checkin_visit_id'] == visit['id']
    quality = domain.completeness(case['id'])
    assert quality['total_images'] == 4
    assert all(r['count'] == 1 and r['missing'] == 0 for r in quality['items'] if r['stage'] == 'work' and 'checkin' in r['role'])
    with services['db'].tx() as con:
        con.execute("UPDATE cases SET workflow_stage='repair' WHERE id=?", (case['id'],))
    board = services['workflow'].board('work', source='local')['items'][0]
    assert board['photo_counts']['survey_checkin'] == board['photo_counts']['work_checkin'] == 2
    assert board['shared_checkin_count'] == 2 and board['work_visit_count'] == 1
    draft = services['platform'].draft(case['id'], 'repair')
    assert draft['files']['buildingFileInfoList'] == [photos[0]['id']]
    assert draft['files']['houseFileInfoList'] == [photos[1]['id']]


def test_second_entry_same_day_has_separate_gaps_and_cannot_reuse_first_checkin(services, case, photo):
    domain = services['domain']
    photos = materials(photo)
    first = record(domain, case, photos)
    later = photo('purple', '下午施工.png', stage='work', role='construction', watermark_time='2026-10-16 15:20')
    second = domain.save_visit(case['id'], entry(visit_key='下午返场', started_at='2026-10-16T15:00', finished_at='2026-10-16T16:00',
                                               same_day_repair=False, photo_ids=[later['id']]))
    quality = domain.completeness(case['id'], 'insurer-v2')
    checks = {r['scope_key']: r for r in quality['items'] if r['photo_group'] == 'work_checkin'}
    assert checks['visit:' + first['id']]['count'] == 2
    assert checks['visit:' + second['id']]['missing'] == 1
    work = {r['scope_key']: r for r in quality['items'] if r['photo_group'] == 'construction'}
    assert work['visit:' + first['id']]['photo_ids'] == [photos[2]['id']]
    assert work['visit:' + second['id']]['photo_ids'] == [later['id']]
    with pytest.raises(DomainError, match='另一次施工'):
        domain.save_visit(case['id'], entry(visit_key='下午返场', revision=second['revision'], shared_checkin_ids=[photos[0]['id']]), second['id'])
    with pytest.raises(DomainError, match='另一次进场'):
        domain.save_visit(case['id'], entry(visit_key='第三次', photo_ids=[photos[2]['id']]))
    assert len(domain.visits(case['id'])) == 2
    assert domain.photos(case_id=case['id'], visit_id=second['id'])['total'] == 1


@pytest.mark.parametrize('change,match', [
    ({'same_day_repair': False}, '标记'),
    ({'visit_date': '2026-10-17', 'started_at': None, 'finished_at': None}, '同一天'),
    ({'finished_at': '2026-10-16T08:00'}, '不能早于'),
    ({'started_at': '2026-10-15T09:00'}, '日期一致'),
])
def test_invalid_reuse_or_time_is_atomic(services, case, photo, change, match):
    domain = services['domain']
    p = photo(stage='survey', role='building_checkin', watermark_time='2026-10-16 09:25')
    with pytest.raises(DomainError, match=match):
        domain.save_visit(case['id'], entry(shared_checkin_ids=[p['id']], **change))
    assert not domain.visits(case['id']) and not services['db'].all('SELECT * FROM visit_shared_checkins')


def test_reuse_rejects_unknown_date_wrong_role_other_case_and_conflicts(services, case, photo):
    domain = services['domain']
    p = photo(stage='survey', role='building_checkin')
    with pytest.raises(DomainError, match='同一天'):
        domain.save_visit(case['id'], entry(shared_checkin_ids=[p['id']]))
    p = domain.update_photo(p['id'], {'watermark_time': '2026-10-16 09:25', 'role': 'problem'})
    with pytest.raises(DomainError, match='共用栏'):
        domain.save_visit(case['id'], entry(shared_checkin_ids=[p['id']]))
    p = domain.update_photo(p['id'], {'role': 'building_checkin'})
    other = domain.save_case({'project_id': case['project_id'], 'title': '另案'})
    with pytest.raises(DomainError, match='必须属于本案'):
        domain.save_visit(other['id'], entry(shared_checkin_ids=[p['id']]))
    with services['db'].tx() as con:
        con.execute("UPDATE photos SET recognition_state='conflict' WHERE id=?", (p['id'],))
    with pytest.raises(DomainError, match='识别冲突'):
        domain.save_visit(case['id'], entry(shared_checkin_ids=[p['id']]))
    assert not domain.visits(case['id'])


def test_edit_release_and_stale_revision_keep_original_classification(services, case, photo, client):
    domain = services['domain']
    photos = materials(photo)
    visit = record(domain, case, photos)
    path = f"/api/cases/{case['id']}/visits/{visit['id']}"
    body = entry(revision=visit['revision'], same_day_repair=False, shared_checkin_ids=[], photo_ids=[])
    saved = client.put(path, json=body)
    assert saved.status_code == 200
    assert saved.json()['counts']['shared_checkin'] == 0
    assert client.put(path, json=body).status_code == 409
    assert domain.photos(case_id=case['id'], photo_group='work_checkin')['total'] == 0
    assert all(domain.photo(p['id'])['visit_id'] is None for p in photos)
    assert domain.photo(photos[0]['id'])['photo_group'] == 'survey_checkin'


def test_photo_identity_change_invalidates_shared_link(services, case, photo):
    domain = services['domain']
    photos = materials(photo)
    visit = record(domain, case, photos)
    domain.update_photo(photos[0]['id'], {'notes': '只改备注'})
    assert len(domain.visits(case['id'])[0]['shared_checkin_ids']) == 2
    domain.update_photo(photos[0]['id'], {'watermark_time': '2026-10-17 09:25'})
    updated = domain.visits(case['id'])[0]
    assert updated['shared_checkin_ids'] == [photos[1]['id']]
    assert updated['counts']['building_checkin'] == 0
    assert domain.photos(case_id=case['id'], photo_group='work_checkin')['total'] == 1


def test_export_preserves_entry_records_gaps_and_one_copy_per_original(services, case, photo):
    domain = services['domain']
    photos = materials(photo)
    first = record(domain, case, photos)
    second = domain.save_visit(case['id'], entry(visit_key='下午返场', started_at=None, finished_at=None, same_day_repair=False, notes='=HYPERLINK("unsafe")'))
    package = services['reporter'].build(uid(), case['id'], [p['id'] for p in photos], 'insurer-v2')
    manifest = package['manifest']
    assert len(manifest['photos']) == 4
    assert {v['id'] for v in manifest['visits']} == {first['id'], second['id']}
    assert any(r['scope_key'] == 'visit:' + second['id'] and r['missing'] for r in manifest['completeness']['items'])
    with zipfile.ZipFile(services['storage'].resolve(package['zip_path'])) as archive:
        for p in manifest['photos']:
            assert hashlib.sha256(archive.read(p['package_path'])).hexdigest() == p['blob_hash']
        workbooks = [f for f in archive.namelist() if f.endswith('.xlsx')]
        workbook = load_workbook(io.BytesIO(archive.read(workbooks[0])))
        sheet = workbook['进场维修记录']
        assert sheet.max_row == 3 and sheet['G3'].data_type == 's'
        assert '楼栋.png' in sheet['I2'].value
    frozen = deepcopy(manifest)
    domain.save_visit(case['id'], entry(visit_key='第一次进场', revision=first['revision'], notes='后改说明'), first['id'])
    assert services['reporter'].package(package['id'])['manifest'] == frozen


@pytest.mark.parametrize('custom', [False, True])
def test_migration_preserves_old_visits_and_custom_requirements(tmp_path, custom):
    from pathlib import Path
    db = Database(tmp_path / 'old.sqlite3')
    original = [{'photo_group': group, 'stage': stage, 'role': role, 'minimum': 1, 'per_visit': False} for group, (stage, role) in GROUP_FIELDS.items()]
    if custom:
        original[2]['minimum'] = 4
    with db.read() as con:
        for script in sorted((Path(__file__).parents[1] / 'idi/migrations').glob('*.sql')):
            version = int(script.name.split('_')[0])
            if version >= 8:
                continue
            con.executescript(script.read_text(encoding='utf-8'))
            con.execute("INSERT INTO schema_version VALUES(?,'2026-10-15')", (version,))
        con.execute("INSERT INTO projects(id,name,short_prefix,created_at,updated_at) VALUES('p','旧项目','OLD','2026-10-15','2026-10-15')")
        con.execute("INSERT INTO cases(id,project_id,title,short_code,created_at,updated_at) VALUES('c','p','旧案','OLD-1-1-01','2026-10-15','2026-10-15')")
        con.execute("INSERT INTO visits VALUES('v','c','work','2026-10-15','旧进场','manual','旧说明','2026-10-15')")
        con.execute("INSERT INTO templates VALUES('insurer-v2','保司基础过程资料','',?,'2026-10-15','2026-10-15')", (dump(original),))
    db.initialize()
    visit = db.one("SELECT * FROM visits WHERE id='v'")
    assert visit['notes'] == '旧说明' and visit['revision'] == 1 and visit['same_day_repair'] == 0
    after = json.loads(db.one("SELECT requirements FROM templates WHERE id='insurer-v2'")['requirements'])
    if custom:
        assert after == original
    else:
        assert {r['photo_group'] for r in after if r['per_visit']} == {'work_checkin', 'construction'}
    db.initialize()
    assert db.one("SELECT * FROM visits WHERE id='v'") == visit
    assert db.all('PRAGMA foreign_key_check') == []
